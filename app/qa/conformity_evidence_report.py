"""Generate an auditable workbook for matrix-to-document evidence candidates.

The report deliberately separates supplier-declared statuses from retrieved
document passages. Neither exact-ID links nor text-similarity links constitute
an independent conformity decision.
"""
from __future__ import annotations

import io
from typing import Any


def generate_evidence_excel(analysis: dict[str, Any]) -> bytes:
    """Build an XLSX review report from PDF/PPTX evidence-linking output."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    requirements = analysis.get("requirements", [])
    summary = analysis.get("summary", {})
    declared = analysis.get("decisionSummary", {}).get("supplierDeclared") or {
        status: sum(item.get("matrixStatus") == status for item in requirements)
        for status in ("OK", "NOK", "NA", "EMPTY")
    }
    pages = "PDF pages" if analysis.get("tdrFile") else "Presentation slides"
    page_count = summary.get("pagesWithExtractedText", summary.get("slidesWithExtractedText", 0))

    wb = Workbook()
    overview = wb.active
    overview.title = "Review Summary"
    header_fill = PatternFill("solid", fgColor="16324F")
    header_font = Font(color="FFFFFF", bold=True)
    section_font = Font(color="16324F", bold=True, size=12)
    wrap = Alignment(vertical="top", wrap_text=True)

    overview["A1"] = "Matrix-to-document evidence review"
    overview["A1"].font = Font(color="16324F", bold=True, size=16)
    overview.merge_cells("A1:D1")
    overview["A3"] = "Matrix file"
    overview["B3"] = analysis.get("matrixFile", "")
    overview["A4"] = "Evidence document"
    overview["B4"] = analysis.get("tdrFile") or analysis.get("presentationFile", "")
    overview["A5"] = "Requirements analyzed"
    overview["B5"] = summary.get("requirements", len(requirements))
    overview["A6"] = "With retrieval candidates"
    overview["B6"] = summary.get("requirementsWithEvidenceCandidates", 0)
    overview["A7"] = "Without text candidates"
    overview["B7"] = summary.get("requirementsWithoutEvidenceCandidates", 0)
    overview["A8"] = f"{pages} containing extracted text"
    overview["B8"] = page_count
    overview["A10"] = "Supplier matrix declarations (not independently verified)"
    overview["A10"].font = section_font
    overview.merge_cells("A10:B10")
    for col, value in enumerate(("Declared status", "Requirement count"), 1):
        cell = overview.cell(row=11, column=col, value=value)
        cell.fill = header_fill
        cell.font = header_font
    for row, status in enumerate(("OK", "NOK", "NA", "EMPTY"), 12):
        overview.cell(row=row, column=1, value=status)
        overview.cell(row=row, column=2, value=declared.get(status, 0))
    overview["A17"] = "Decision policy"
    overview["A17"].font = section_font
    overview["A18"] = (
        "This report links requirements to text passages for human review. "
        "Supplier statuses remain declarations from the matrix; this analysis "
        "does not independently classify conformity. A missing text candidate "
        "is not evidence of non-conformity."
    )
    overview.merge_cells("A18:D19")
    overview["A18"].alignment = wrap
    limitations = analysis.get("limitations", [])
    overview["A21"] = "Known limitations / reviewer checks"
    overview["A21"].font = section_font
    for row, text in enumerate(limitations, 22):
        overview.cell(row=row, column=1, value=text).alignment = wrap
        overview.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    overview.column_dimensions["A"].width = 42
    overview.column_dimensions["B"].width = 34
    overview.column_dimensions["C"].width = 22
    overview.column_dimensions["D"].width = 22
    overview.freeze_panes = "A11"

    requirements_sheet = wb.create_sheet("Requirements")
    req_headers = [
        "Matrix row", "Requirement ID", "Reference", "Requirement description",
        "Supplier declared status", "Raw matrix status", "Supplier comment",
        "Evidence review state", "Candidate count",
    ]
    for column, label in enumerate(req_headers, 1):
        cell = requirements_sheet.cell(row=1, column=column, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap
    candidate_sheet = wb.create_sheet("Evidence Candidates")
    evidence_headers = [
        "Matrix row", "Requirement ID", "Supplier declared status", "Match type",
        "Retrieval score (not confidence)", "Source document", "Page / slide",
        "Shape", "Evidence excerpt",
    ]
    for column, label in enumerate(evidence_headers, 1):
        cell = candidate_sheet.cell(row=1, column=column, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap

    candidate_row = 2
    for row_number, item in enumerate(requirements, 2):
        candidates = item.get("evidence", []) or []
        values = [
            item.get("matrixRowNumber", item.get("rowIndex", "")),
            item.get("reqId", ""), item.get("reference", ""),
            item.get("description", ""), item.get("matrixStatus", ""),
            item.get("matrixStatusRaw", ""), item.get("supplierComment", item.get("comment", "")),
            item.get("reviewStatus", ""), len(candidates),
        ]
        for column, value in enumerate(values, 1):
            requirements_sheet.cell(row=row_number, column=column, value=value).alignment = wrap

        for candidate in candidates:
            citation = candidate.get("page_number", candidate.get("slide_number", ""))
            candidate_values = [
                item.get("matrixRowNumber", item.get("rowIndex", "")),
                item.get("reqId", ""), item.get("matrixStatus", ""),
                candidate.get("matchType", ""), candidate.get("matchScore", ""),
                candidate.get("file_name", ""), citation,
                candidate.get("shape_name", ""), candidate.get("text", ""),
            ]
            for column, value in enumerate(candidate_values, 1):
                candidate_sheet.cell(row=candidate_row, column=column, value=value).alignment = wrap
            candidate_row += 1

    req_widths = [12, 20, 32, 70, 24, 20, 60, 32, 16]
    for column, width in enumerate(req_widths, 1):
        requirements_sheet.column_dimensions[get_column_letter(column)].width = width
    candidate_widths = [12, 20, 24, 24, 32, 45, 14, 28, 85]
    for column, width in enumerate(candidate_widths, 1):
        candidate_sheet.column_dimensions[get_column_letter(column)].width = width
    requirements_sheet.freeze_panes = "A2"
    candidate_sheet.freeze_panes = "A2"
    requirements_sheet.auto_filter.ref = f"A1:I{max(1, len(requirements) + 1)}"
    candidate_sheet.auto_filter.ref = f"A1:I{max(1, candidate_row - 1)}"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()