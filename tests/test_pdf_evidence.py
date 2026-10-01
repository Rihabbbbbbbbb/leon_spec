from pathlib import Path
from types import SimpleNamespace

import pytest

from app.qa.pdf_evidence import (
    PageEvidence,
    analyze_matrix_against_pdf,
    extract_pdf_evidence,
    match_requirement_evidence,
)


def test_extract_pdf_evidence_keeps_empty_pages_out_of_results(tmp_path):
    from PyPDF2 import PdfWriter

    path = tmp_path / "tdr.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as stream:
        writer.write(stream)

    assert extract_pdf_evidence(path) == []


def test_pdf_extractor_validates_extension_and_missing_file(tmp_path):
    with pytest.raises(ValueError, match="Only .pdf"):
        extract_pdf_evidence(tmp_path / "tdr.docx")
    with pytest.raises(FileNotFoundError):
        extract_pdf_evidence(tmp_path / "missing.pdf")


def test_match_prefers_exact_requirement_id_over_lexical_candidates():
    requirement = {
        "reqId": "REQ-12345",
        "reference": "GEN-PHYS-001",
        "description": "The display shall operate at 12 volts under nominal conditions.",
    }
    pages = [
        PageEvidence("tdr.pdf", 2, "Display operation nominal voltage is twelve volts."),
        PageEvidence(
            "tdr.pdf", 9, "REQ-12345: Display shall operate at 12 V nominal.",
            evidence_ids=["REQ-12345"],
        ),
    ]

    matches = match_requirement_evidence(requirement, pages)

    assert matches[0]["page_number"] == 9
    assert matches[0]["matchType"] == "requirement_id"
    assert matches[0]["matchScore"] == 1.0
    assert all(match["retrievalOnly"] is True for match in matches)


def test_weak_or_absent_text_is_not_claimed_as_evidence():
    requirement = {"reqId": "REQ-12345", "description": "Optical attenuation shall remain below 0.5 dB."}
    pages = [PageEvidence("tdr.pdf", 1, "Supplier company profile and contact details.")]

    assert match_requirement_evidence(requirement, pages) == []


def test_matrix_analysis_never_generates_compliance_verdict(tmp_path, monkeypatch):
    from app.qa import pdf_evidence

    monkeypatch.setattr(
        pdf_evidence,
        "extract_pdf_evidence",
        lambda _: [PageEvidence(
            "tdr.pdf", 4, "REQ-12345: measured value = 11.8 V",
            evidence_ids=["REQ-12345"],
        )],
    )
    analysis = SimpleNamespace(
        file_name="supplier.xlsm",
        items=[SimpleNamespace(
            is_requirement=True,
            row_index=7,
            req_id="REQ-12345",
            reference="",
            description="Input voltage shall be 12 V nominal.",
            conformity_category="OK",
            conformity_raw="OK",
            comment="Test report referenced.",
        )],
    )

    result = analyze_matrix_against_pdf(analysis, tmp_path / "tdr.pdf")

    row = result["requirements"][0]
    assert row["matrixRowNumber"] == 8
    assert row["matrixStatus"] == "OK"  # supplier declaration, not the system verdict
    assert row["evidence"][0]["page_number"] == 4
    assert result["summary"]["complianceVerdictsMade"] == 0
    assert "verdict" not in row
    assert "No text candidate found does not mean the supplier failed" in " ".join(result["limitations"])


def test_pdf_extractor_repairs_spaced_requirement_id_and_preserves_page(monkeypatch, tmp_path):
    from app.qa import pdf_evidence

    class FakePage:
        def extract_text(self):
            return "REQ - 0308287\nConsommation électrique nominale"

    class FakeReader:
        pages = [FakePage()]

    monkeypatch.setattr(pdf_evidence, "PdfReader", lambda _: FakeReader())
    path = tmp_path / "spaced-id.pdf"
    path.write_bytes(b"test")

    evidence = extract_pdf_evidence(path)

    assert len(evidence) == 1
    assert evidence[0].page_number == 1
    assert "REQ-0308287" in evidence[0].evidence_ids
    assert "electrique" in pdf_evidence._tokens(evidence[0].text)


def test_spaced_pdf_requirement_id_matches_matrix_id(monkeypatch, tmp_path):
    from app.qa import pdf_evidence

    class FakePage:
        def extract_text(self):
            return "REQ - 0308287\nNominal current consumption: 85 mA"

    class FakeReader:
        pages = [FakePage()]

    monkeypatch.setattr(pdf_evidence, "PdfReader", lambda _: FakeReader())
    path = tmp_path / "spaced-id.pdf"
    path.write_bytes(b"test")
    evidence = extract_pdf_evidence(path)

    matches = match_requirement_evidence(
        {"reqId": "REQ-0308287", "description": "Nominal current consumption shall not exceed 100 mA."},
        evidence,
    )

    assert matches
    assert matches[0]["matchType"] == "requirement_id"
    assert matches[0]["page_number"] == 1


def test_evidence_report_keeps_unmatched_rows_and_supplier_declarations():
    from io import BytesIO

    from openpyxl import load_workbook

    from app.qa.conformity_evidence_report import generate_evidence_excel

    result = {
        "matrixFile": "supplier.xlsm",
        "tdrFile": "tdr.pdf",
        "summary": {"requirements": 2, "requirementsWithEvidenceCandidates": 1,
                    "requirementsWithoutEvidenceCandidates": 1, "pagesWithExtractedText": 2},
        "decisionSummary": {"supplierDeclared": {"OK": 1, "NOK": 1, "NA": 0, "EMPTY": 0}},
        "requirements": [
            {"matrixRowNumber": 8, "reqId": "REQ-1", "description": "A requirement",
             "matrixStatus": "OK", "matrixStatusRaw": "OK", "supplierComment": "Claim",
             "evidence": [{"matchType": "requirement_id", "matchScore": 1.0,
                           "file_name": "tdr.pdf", "page_number": 3, "text": "REQ-1 result"}]},
            {"matrixRowNumber": 9, "reqId": "REQ-2", "description": "Another requirement",
             "matrixStatus": "NOK", "matrixStatusRaw": "NOK", "supplierComment": "",
             "evidence": []},
        ],
    }

    workbook = load_workbook(BytesIO(generate_evidence_excel(result)), data_only=True)

    assert workbook.sheetnames == ["Review Summary", "Requirements", "Evidence Candidates"]
    assert workbook["Review Summary"]["B12"].value == 1
    assert workbook["Requirements"]["A3"].value == 9
    assert workbook["Requirements"]["H3"].value in (None, "")
    assert workbook["Evidence Candidates"]["G2"].value == 3
    assert workbook["Evidence Candidates"]["I2"].value == "REQ-1 result"
