from __future__ import annotations

import base64
from io import BytesIO
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.qa import conformity_analyzer
from app.conformity_server import app


client = TestClient(app)


def _matrix_bytes(status: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Requirement", "Description", "Conformity FNR", "Comments FNR"])
    sheet.append(["REQ-12345", "The indicator shall illuminate within 100 ms.", status, "Demo only."])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_compare_request_is_self_contained_and_generates_delta_report():
    response = client.post(
        "/api/aeris-conformity-compare",
        files=[
            (
                "files",
                (
                    "AERIS_Demo_V1.xlsx",
                    _matrix_bytes("OK"),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            (
                "files",
                (
                    "AERIS_Demo_V2.xlsx",
                    _matrix_bytes("NOK"),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
        ],
    )

    assert response.status_code == 200, response.text
    result = response.json()
    assert result["totalChanges"] == 1
    assert len(result["steps"]) == 1
    assert result["steps"][0]["statusChanges"] == 1
    report = load_workbook(
        BytesIO(base64.b64decode(result["reportExcel"])), data_only=True
    )
    assert report.sheetnames == [
        "Delta Overview",
        "Status Changes",
        "New & Removed",
        "Comment Changes",
    ]


def test_compare_requires_two_supported_nonempty_files():
    one_file = client.post(
        "/api/aeris-conformity-compare",
        files=[
            (
                "files",
                (
                    "AERIS_Demo.xlsx",
                    _matrix_bytes("OK"),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            )
        ],
    )
    unsupported = client.post(
        "/api/aeris-conformity-compare",
        files=[
            ("files", ("AERIS_Demo.csv", b"not a matrix", "text/csv")),
            ("files", ("AERIS_Demo_V2.xlsx", _matrix_bytes("NOK"), "application/octet-stream")),
        ],
    )
    empty = client.post(
        "/api/aeris-conformity-compare",
        files=[
            ("files", ("AERIS_Demo.xlsx", b"", "application/octet-stream")),
            ("files", ("AERIS_Demo_V2.xlsx", _matrix_bytes("NOK"), "application/octet-stream")),
        ],
    )

    assert one_file.status_code == 400
    assert unsupported.status_code == 400
    assert empty.status_code == 400


def test_compare_builds_compatible_report_when_analyzer_has_no_steps(monkeypatch):
    monkeypatch.setattr(
        conformity_analyzer,
        "compare_matrices",
        lambda paths, names: SimpleNamespace(),
    )
    monkeypatch.setattr(
        conformity_analyzer,
        "comparison_to_dict",
        lambda comparison: {
            "statusChanges": [
                {
                    "step": 1,
                    "reqId": "REQ-12345",
                    "from": "OK",
                    "to": "NOK",
                    "changeType": "regressed",
                }
            ],
            "totalChanges": 1,
            "reportText": "Fictional demo comparison.",
        },
    )
    response = client.post(
        "/api/aeris-conformity-compare",
        files=[
            ("files", ("AERIS_Demo_V1.xlsx", _matrix_bytes("OK"), "application/octet-stream")),
            ("files", ("AERIS_Demo_V2.xlsx", _matrix_bytes("NOK"), "application/octet-stream")),
        ],
    )

    assert response.status_code == 200, response.text
    result = response.json()
    assert result["steps"] == []
    report = load_workbook(
        BytesIO(base64.b64decode(result["reportExcel"])), data_only=True
    )
    assert report.sheetnames == ["Delta Overview", "Status Changes", "Review Notes"]
    assert report["Status Changes"]["B2"].value == "REQ-12345"
