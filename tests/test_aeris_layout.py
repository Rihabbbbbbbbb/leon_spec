"""Presentation-independent evidence extraction and grounded comparison."""
from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.qa.aeris_crosscheck import run_crosscheck
from app.qa.aeris_evidence import parse_evidence_bytes
from app.qa.aeris_constraints import extract_measurements


def _matrix(tmp_path, description="Display refresh rate >=40 Hz"):
    path = tmp_path / "matrix.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append([
        "Numéro de l'exigence", "Référence", "", "", "",
        "Libellé", "Conformité FNR", "Commentaires FNR",
    ])
    sheet.append(["REQ-1234", "Display", "", "", "", description, "OK", ""])
    book.save(path)
    return path


def _pptx():
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    image = io.BytesIO()
    Image.new("RGB", (640, 120), "white").save(image, format="PNG")
    image.seek(0)
    presentation = Presentation()
    for i in range(3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(6.9), Inches(8), Inches(0.3)
        ).text = f"{i + 1} © 2026 Example Co. All rights reserved"
        if i == 1:
            table = slide.shapes.add_table(
                3, 2, Inches(1), Inches(1), Inches(7), Inches(2)
            ).table
            for row, cells in enumerate([
                ("Measurement", "Result"),
                ("PWM frequency", "40 Hz"),
                ("Display refresh rate", "45 Hz"),
            ]):
                for col, text in enumerate(cells):
                    table.cell(row, col).text = text
            slide.shapes.add_picture(image, Inches(2), Inches(4), width=Inches(4))
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def test_pptx_rows_not_conflated_and_footer_removed(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "app.qa.aeris_evidence._ocr_image",
        lambda image, warnings: "Display refresh rate measured 45 Hz",
    )
    content = _pptx()
    parsed = parse_evidence_bytes("vendor.pptx", content)
    assert parsed.parse_error == ""
    assert parsed.page_count == 1
    assert {c.kind for c in parsed.chunks} == {"table_row", "image"}
    assert not any("rights reserved" in c.text for c in parsed.chunks)
    assert any("Measurement: Display refresh rate" in c.text for c in parsed.chunks)
    report = run_crosscheck(str(_matrix(tmp_path)), [("vendor.pptx", content)])
    item = report.items[0]
    assert item.final_status == "CONFORME"
    assert "45" in item.supplier_result
    assert "40" not in item.supplier_result
    assert item.evidence_location == "Slide 2"


def test_frequency_only_distractor_is_not_evidence(tmp_path):
    path = _matrix(tmp_path)
    text = b"Slide 1\nPWM frequency measured 40 Hz\n\nSlide 2\nBattery voltage 12 V"
    report = run_crosscheck(str(path), [("alternate.txt", text)])
    assert report.items[0].final_status in {"MANQUANT", "PREUVE_INSUFFISANTE"}
    assert not report.items[0].condition_verdicts


def test_two_frequency_properties_in_one_paragraph(tmp_path):
    values = extract_measurements(
        "PWM frequency measured 40 Hz; display refresh rate measured 45 Hz"
    )
    assert [(m.quantity, m.value) for m in values] == [
        ("pwm_frequency", 40), ("refresh_frequency", 45),
    ]
    matrix = _matrix(tmp_path)
    text = b"Slide 2\nPWM frequency measured 40 Hz; display refresh rate measured 45 Hz"
    item = run_crosscheck(str(matrix), [("another.txt", text)]).items[0]
    assert item.final_status == "CONFORME"
    assert "45" in item.supplier_result
    assert "40" not in item.supplier_result


def test_pdf_table_image_and_page_footer(tmp_path, monkeypatch):
    pytest.importorskip("pdfplumber")
    monkeypatch.setattr(
        "app.qa.aeris_evidence._ocr_image",
        lambda image, warnings: "Display refresh rate measured 45 Hz",
    )
    from fpdf import FPDF
    from PIL import Image
    image = io.BytesIO()
    Image.new("RGB", (640, 120), "white").save(image, format="PNG")
    pdf = FPDF()
    for i in range(3):
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        pdf.set_y(270)
        pdf.cell(text=f"{i + 1} (C) 2026 Example Co. All rights reserved")
        if i == 1:
            pdf.set_y(30)
            pdf.cell(text="PWM frequency measured 40 Hz")
            pdf.image(image, x=20, y=55, w=150)
    content = bytes(pdf.output())
    doc = parse_evidence_bytes("other-layout.pdf", content)
    assert not doc.parse_error
    assert any(c.kind == "image" for c in doc.chunks)
    assert not any("rights reserved" in c.text for c in doc.chunks)
    report = run_crosscheck(str(_matrix(tmp_path)), [("other-layout.pdf", content)])
    assert report.items[0].final_status == "CONFORME"
    assert "45" in report.items[0].supplier_result


def test_unsupported_legacy_ppt_has_explicit_error():
    doc = parse_evidence_bytes("old.ppt", b"not a PPTX")
    assert not doc.chunks
    assert "not supported" in doc.parse_error


def test_ocr_unavailable_is_reported(monkeypatch):
    from app.qa import aeris_evidence
    import pytesseract
    from PIL import Image
    image = io.BytesIO()
    Image.new("RGB", (640, 120), "white").save(image, format="PNG")

    def no_ocr(_):
        raise pytesseract.TesseractNotFoundError()

    monkeypatch.setattr(pytesseract, "image_to_string", no_ocr)
    warnings = []
    assert aeris_evidence._ocr_image(image.getvalue(), warnings) == ""
    assert any("OCR unavailable" in message for message in warnings)
    doc = parse_evidence_bytes("ocr-required.pptx", _pptx())
    assert any("OCR unavailable" in message for message in doc.warnings)


def test_http_rejects_legacy_ppt(tmp_path):
    from fastapi.testclient import TestClient
    from app.conformity_server import app

    with _matrix(tmp_path).open("rb") as matrix:
        response = TestClient(app).post(
            "/api/aeris-crosscheck",
            files=[
                ("matrix", ("matrix.xlsx", matrix, "application/octet-stream")),
                ("evidence", ("old.ppt", b"legacy data", "application/octet-stream")),
            ],
        )
    assert response.status_code == 400
    assert ".pptx" in response.json()["detail"]


def test_http_failed_extraction_is_not_a_successful_all_missing_report(tmp_path):
    from fastapi.testclient import TestClient
    from app.conformity_server import app

    with _matrix(tmp_path).open("rb") as matrix:
        response = TestClient(app).post(
            "/api/aeris-crosscheck",
            files=[
                ("matrix", ("matrix.xlsx", matrix, "application/octet-stream")),
                ("evidence", ("broken.pdf", b"not a PDF", "application/pdf")),
            ],
        )
    assert response.status_code == 422
    assert "No usable TDR evidence" in response.json()["detail"]
    assert "items" not in response.json()
