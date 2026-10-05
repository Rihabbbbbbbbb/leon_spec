"""Adversarial accuracy regressions for deterministic matrix/TDR decisions."""
from dataclasses import replace
import io

import pytest
from openpyxl import load_workbook

from app.qa.aeris_constraints import (
    compare_constraint, extract_constraints, extract_measurements, summarize_verdicts,
)
from app.qa.aeris_crosscheck import _check_one
from app.qa.aeris_evidence import EvidenceChunk, parse_evidence_bytes
from app.qa.aeris_report import generate_aeris_excel
from app.qa.aeris_statements import extract_restated_targets, restated_vs_requirement
from app.qa.conformity_analyzer import ConformityItem


def check(description, *texts):
    req = ConformityItem(
        row_index=1, req_id="REQ-1234", description=description,
        conformity_category="OK",
    )
    chunks = [
        EvidenceChunk(f"vendor-{i}.txt", f"Page {i}", text, i)
        for i, text in enumerate(texts, 1)
    ]
    return _check_one(req, chunks)[0]


def test_two_sided_limits_keep_their_own_operators():
    cs = extract_constraints("Voltage >=9V and <=16V")
    assert [(c.value, c.operator) for c in cs] == [(9, "ge"), (16, "le")]


@pytest.mark.parametrize("wording,operator", [
    ("shall be greater than", "gt"),
    ("must be less than", "lt"),
    ("shall be greater than or equal to", "ge"),
    ("shall be less than or equal to", "le"),
    ("shall not exceed", "le"),
    ("at least", "ge"),
])
def test_worded_inequalities_preserve_strictness(wording, operator):
    constraints = extract_constraints(f"Current {wording} 100mA")
    assert len(constraints) == 1
    assert constraints[0].operator == operator


def test_missing_second_constraint_never_establishes_compliance():
    item = check("Current <=100mA; display temperature <38°C", "Measured current 80mA")
    assert item.final_status == "PREUVE_INSUFFISANTE"
    assert any(v["status"] == "INCOMPARABLE" for v in item.condition_verdicts)
    assert "38" in item.target


def test_missing_constraint_does_not_hide_a_proven_failure():
    item = check("Current <=100mA; display temperature <38°C", "Measured current 120mA")
    assert item.final_status == "NON_CONFORME"


def test_supplier_restated_limit_is_not_an_authoritative_customer_target():
    req = ConformityItem(
        row_index=1, req_id="REQ-1234", description="Provide a current consumption report",
        conformity_category="OK", comment="Current target <=100mA",
    )
    item = _check_one(req, [EvidenceChunk(
        "vendor.txt", "Page 1", "REQ-1234 Current measured 80mA", 1,
    )])[0]
    assert item.final_status == "PREUVE_INSUFFISANTE"
    assert item.target == ""


def test_multiple_requirement_results_in_one_unstructured_chunk_are_not_conflated():
    item = check(
        "Current <=100mA",
        "REQ-1234 current measured 80mA\nREQ-9999 current measured 200mA",
    )
    assert item.final_status in {"MANQUANT", "PREUVE_INSUFFISANTE"}
    assert not item.condition_verdicts


def test_other_requirement_id_cannot_prove_this_requirement():
    item = check("Current <=100mA", "REQ-9999 Current measured 80mA")
    assert item.final_status == "MANQUANT"


def test_voltage_operating_point_is_not_an_acceptance_limit():
    cs = extract_constraints("Current <=100mA at 13.5V")
    assert len(cs) == 1
    assert "13.5" in cs[0].condition
    assert check(
        "Current <=100mA at 13.5V", "Current measured 80mA at 9V",
    ).final_status != "CONFORME"


def test_unspecified_voltage_does_not_prove_required_operating_point():
    assert check(
        "Current <=100mA at 13.5V", "Current measured 80mA",
    ).final_status == "PREUVE_INSUFFISANTE"


@pytest.mark.parametrize("value", [2000, 2026, 100000])
def test_year_like_and_long_numeric_measurements_are_not_identifiers(value):
    cs = extract_constraints(f"Current <={value}mA")
    ms = extract_measurements(f"Current measured {value}mA")
    assert len(cs) == len(ms) == 1
    assert cs[0].value == ms[0].value == value


def test_operating_condition_dimensions_do_not_match_by_number_alone():
    from app.qa.aeris_constraints import conditions_compatible
    assert not conditions_compatible("@32°C", "@V32°")
    assert not conditions_compatible("at 32V", "@32°C")


@pytest.mark.parametrize("text,expected", [
    ("Not compliant", "FAIL"),
    ("Not OK", "FAIL"),
    ("Non-compliant", "FAIL"),
    ("Not passed", "FAIL"),
    ("No failures. Result OK", "PASS"),
    ("Without deviations. Result OK", "PASS"),
])
def test_negated_declarations_are_not_reversed(text, expected):
    from app.qa.aeris_contradictions import tdr_polarity
    assert tdr_polarity(text) == expected


def test_explicit_equality_is_not_a_default_temperature_upper_limit():
    cs = extract_constraints("Display temperature =25°C")
    assert cs[0].operator == "eq"
    vs = compare_constraint(cs[0], extract_measurements("Display temperature measured 24°C"))
    assert summarize_verdicts(vs) == "NON_CONFORME"


def test_equality_has_no_undocumented_one_percent_tolerance():
    c = replace(extract_constraints("Voltage <=100V")[0], operator="eq")
    vs = compare_constraint(c, extract_measurements("Voltage measured 100.5V"))
    assert summarize_verdicts(vs) == "NON_CONFORME"


def test_max_label_does_not_hide_a_higher_measured_result():
    item = check("Current <=100mA", "Current measured 120mA; Max 90mA")
    assert item.final_status != "CONFORME"


def test_restated_operator_mismatch_is_a_wrong_target():
    constraints = extract_constraints("Current <=100mA")
    restated = extract_restated_targets("Current target >=100mA")
    assert restated_vs_requirement(restated, constraints)[1] == "WRONG_TARGET"


def test_restated_second_target_is_checked_too():
    constraints = extract_constraints("Current <=100mA; display temperature <38°C")
    restated = extract_restated_targets("Current target <=100mA; display temperature target <45°C")
    assert restated_vs_requirement(restated, constraints)[1] == "WRONG_TARGET"


def test_all_exact_rows_contribute_even_after_six_passes():
    texts = [f"REQ-1234 Current measured {70 + i}mA" for i in range(7)]
    texts.append("REQ-1234 Current measured 120mA")
    assert check("Current <=100mA", *texts).final_status != "CONFORME"


def test_supplier_table_does_not_hide_explicit_failure_on_another_page():
    req = ConformityItem(
        row_index=1, req_id="REQ-1234", description="Current <=100mA",
        conformity_category="OK",
    )
    chunks = [
        EvidenceChunk(
            "vendor.pdf", "Page 1", "REQ-1234 Description: Current <=100mA | Result: 80mA",
            1, "table_row", "Current measured 80mA",
        ),
        EvidenceChunk("vendor.pdf", "Page 2", "REQ-1234 Current measured 120mA", 2),
    ]
    item = _check_one(req, chunks)[0]
    assert item.final_status != "CONFORME"
    assert any(v["evidence_location"] == "Page 2" for v in item.condition_verdicts)


def test_explicit_declarations_on_different_pages_cannot_hide_each_other():
    item = check(
        "EMC immunity shall be compliant",
        "REQ-1234 EMC immunity tested. Result OK",
        "REQ-1234 EMC immunity tested. Result NOK",
    )
    assert item.tdr_polarity == "MIXED"
    assert item.contradiction_type == "TDR_INTERNAL_CONFLICT"


def test_equality_considers_every_result_even_if_one_matches():
    item = check("Voltage =100V", "Voltage measured 99.5V and 100V")
    assert item.final_status != "CONFORME"


def test_exact_failure_is_not_hidden_by_an_overlapping_bound():
    item = check("Contrast >=400:1", "Contrast >380:1 and measured 390:1")
    assert item.final_status == "NON_CONFORME"


def test_condition_verdict_points_to_actual_source_document():
    item = check(
        "Current <=100mA",
        "REQ-1234 Current test report available",
        "REQ-1234 Current measured 120mA",
    )
    v = item.condition_verdicts[0]
    assert v["evidence_file"] == "vendor-2.txt"
    assert v["evidence_location"] == "Page 2"
    assert "120" in v["evidence_excerpt"]


def test_docx_table_copied_spec_is_not_supplier_measurement():
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for row, cells in enumerate([
        ["Req ID", "Description", "Supplier proposal"],
        ["REQ-1234", "Current <=100mA", "Measured current 120mA"],
    ]):
        for col, text in enumerate(cells):
            table.cell(row, col).text = text
    stream = io.BytesIO()
    doc.save(stream)
    parsed = parse_evidence_bytes("vendor.docx", stream.getvalue())
    assert not parsed.parse_error
    rows = [c for c in parsed.chunks if c.kind == "table_row"]
    assert len(rows) == 1
    assert "100" not in rows[0].measurement_text


def test_corrupt_docx_exposes_parse_error():
    parsed = parse_evidence_bytes("broken.docx", b"not a zip")
    assert parsed.parse_error
    assert not parsed.chunks


def test_invalid_text_encoding_does_not_silently_replace_engineering_symbols():
    parsed = parse_evidence_bytes("broken.txt", b"Measured current \xff80mA")
    assert parsed.parse_error
    assert not parsed.chunks


def test_aeris_route_runs_analysis_off_the_event_loop(tmp_path, monkeypatch):
    import asyncio
    import threading
    from fastapi import UploadFile
    from app.qa import aeris_crosscheck, route
    from tests.test_aeris_robustness import write_matrix
    path = write_matrix(tmp_path / "matrix.xlsx", [
        ("REQ-1234", "EE", "Current <=100mA", "OK", ""),
    ])
    main_thread = threading.get_ident()
    analysis_threads = []
    original = aeris_crosscheck.run_crosscheck
    def observed(*args, **kwargs):
        analysis_threads.append(threading.get_ident())
        return original(*args, **kwargs)
    monkeypatch.setattr(aeris_crosscheck, "run_crosscheck", observed)
    payload = asyncio.run(route.aeris_crosscheck(
        matrix=UploadFile(filename="matrix.xlsx", file=io.BytesIO(path.read_bytes())),
        evidence=[UploadFile(
            filename="tdr.txt", file=io.BytesIO(b"REQ-1234 Current measured 80mA"),
        )],
    ))
    assert analysis_threads and analysis_threads[0] != main_thread
    assert payload["items"][0]["final_status"] == "CONFORME"
    assert payload["reportExcel"]


def test_whitespace_only_evidence_is_not_a_chunk():
    assert not parse_evidence_bytes("empty.txt", b" \n\n ").chunks
    assert not parse_evidence_bytes("footer.txt", b"Copyright Example").chunks


def test_excel_preserves_gap_and_source():
    report = {
        "incoherences": [{
            "req_id": "REQ-1234", "n": 1, "écart": "+20 mA", "ou": "vendor.txt: Page 2",
        }],
        "items": [], "summary": {},
    }
    wb = load_workbook(io.BytesIO(generate_aeris_excel(report)))
    assert wb["Incohérences"]["J6"].value == "+20 mA"
    assert wb["Incohérences"]["H6"].value == "vendor.txt: Page 2"


def test_all_verdicts_and_sources_survive_excel_export():
    item = check("Current <=100mA", "REQ-1234 Current measured 120mA")
    from dataclasses import asdict
    wb = load_workbook(io.BytesIO(generate_aeris_excel({
        "items": [asdict(item)], "summary": {}, "incoherences": [],
    })))
    detail = wb["Détail complet"]
    assert "vendor-1.txt" in detail["L2"].value
    assert "NON_CONFORME" in detail["M2"].value
    assert detail["N2"].value == item.rationale


@pytest.mark.parametrize("operator", ["le", "lt", "ge", "gt", "eq"])
@pytest.mark.parametrize("bound", ["le", "lt", "ge", "gt"])
@pytest.mark.parametrize("value", [90, 100, 110])
def test_bound_decisions_match_an_independent_interval_oracle(operator, bound, value):
    """60 bound/target combinations, including open and closed endpoints."""
    c = replace(extract_constraints("Current <=100mA")[0], operator=operator)
    m = replace(extract_measurements(f"Current measured {value}mA")[0], bound=bound)
    from fractions import Fraction
    candidates = [Fraction(i, 2) for i in range(-400, 401)]
    def accepts(op, observed, target):
        return {
            "le": observed <= target, "lt": observed < target,
            "ge": observed >= target, "gt": observed > target, "eq": observed == target,
        }[op]
    possibilities = [n for n in candidates if accepts(bound, n, value)]
    outcomes = {accepts(operator, n, 100) for n in possibilities}
    expected = (
        "CONFORME" if outcomes == {True}
        else "NON_CONFORME" if outcomes == {False}
        else "INCOMPARABLE"
    )
    assert compare_constraint(c, [m])[0].status == expected


@pytest.mark.parametrize("description", [
    "Operating temperature from -40°C to 85°C",
    "Mechanical gap 10 +/- 0.5 mm",
    "Current <=100mA or power <=500mW",
])
def test_unsupported_range_tolerance_or_alternative_is_not_certified(description):
    item = check(description, "Current measured 80mA; temperature 20°C; gap 0.3 mm")
    assert item.final_status != "CONFORME"
