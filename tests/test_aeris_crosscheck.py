"""
AERIS matrix ↔ TDR engine — unit + pipeline tests.

Fixtures reconstruct the three Tianma-style cases described in the
product conversation (current, LCF attenuation, TFT contrast) plus the
CPU-load / display-temperature reasoning examples. The real supplier
files are not in this repository; these fixtures use the same numbers
and wording so the deterministic engine can be proven.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

from app.qa.aeris_constraints import (
    compare_constraint,
    extract_constraints,
    extract_measurements,
    summarize_verdicts,
)
from app.qa.aeris_contradictions import classify_contradiction, tdr_polarity
from app.qa.aeris_crosscheck import crosscheck_analysis, report_to_dict, run_crosscheck
from app.qa.aeris_evidence import parse_evidence_bytes
from app.qa.aeris_report import generate_aeris_excel
from app.qa.aeris_statements import (
    extract_restated_targets,
    matrix_polarity,
    values_relation,
)
from app.qa.conformity_analyzer import classify_conformity, extract_conformity_data


# ── Isolated constraint / measurement parsing ─────────────────────

class TestConstraintExtraction:

    def test_current_upper_bound(self):
        cs = extract_constraints(
            "Current consumption in Reduced Consumption Mode ≤100mA"
        )
        assert cs
        c = cs[0]
        assert c.unit_family == "current"
        assert c.value == 100
        assert c.operator == "le"
        assert "reduc" in c.condition.lower()

    def test_lcf_percent_at_angle(self):
        cs = extract_constraints("LCF Attenuation >95% at V=32°")
        assert cs
        c = next(x for x in cs if x.unit_family == "percent")
        assert c.value == 95
        assert c.operator == "gt"
        assert "32" in c.condition

    def test_contrast_ratio(self):
        cs = extract_constraints("TFT contrast ratio ≥400:1")
        assert cs
        c = cs[0]
        assert c.unit_family == "ratio"
        assert c.value == 400
        assert c.operator == "ge"

    def test_skips_req_id_numbers(self):
        cs = extract_constraints("REQ-0308287 shall keep the display on")
        assert cs == []

    def test_temperature_upper_bound(self):
        cs = extract_constraints("Display temperature <38°C")
        assert cs
        assert cs[0].unit_family == "temperature"
        assert cs[0].value == 38
        assert cs[0].operator == "lt"


class TestMeasurementExtraction:

    def test_typ_and_max_current(self):
        ms = extract_measurements(
            "Reduced mode current Typ 119.9mA Max 192.7mA NOK Deviation"
        )
        currents = [m for m in ms if m.unit_family == "current"]
        assert {round(m.value, 1) for m in currents} >= {119.9, 192.7}
        assert any(m.qualifier == "typ" for m in currents)
        assert any(m.qualifier == "max" for m in currents)

    def test_lcf_two_angles(self):
        ms = extract_measurements(
            "lum attenuates >85% @V=32°\nlum attenuates >95% @V=44°"
        )
        pct = [m for m in ms if m.unit_family == "percent"]
        assert any(abs(m.value - 85) < 0.01 and "32" in m.condition for m in pct)
        assert any(abs(m.value - 95) < 0.01 and "44" in m.condition for m in pct)

    def test_contrast_by_temperature(self):
        ms = extract_measurements("TFT contrast 500 @25°C  410 @70°C  380 @85°C")
        ratios = [m for m in ms if m.unit_family == "ratio"]
        assert {round(m.value) for m in ratios} >= {500, 410, 380}


class TestNumericCompare:

    def test_current_uses_worst_max(self):
        cs = extract_constraints("Reduced Consumption Mode ≤100mA")
        ms = extract_measurements("Typ 119.9mA Max 192.7mA")
        vs = compare_constraint(cs[0], ms)
        assert vs
        assert summarize_verdicts(vs) == "NON_CONFORME"
        # Worst-case max 192.7 is the compared value for an upper bound.
        assert any(abs(v.gap - 92.7) < 0.2 or abs(v.gap - 19.9) < 0.2 for v in vs)

    def test_lcf_lower_bound_at_32deg_is_not_an_exact_failure(self):
        cs = extract_constraints("LCF Attenuation >95% at V=32°")
        ms = extract_measurements("lum attenuates >85% @V=32°  lum attenuates >95% @V=44°")
        vs = compare_constraint(cs[0], ms)
        assert summarize_verdicts(vs) == "PREUVE_INSUFFISANTE"
        assert all(v.status == "INCOMPARABLE" and v.gap is None for v in vs)

    def test_contrast_partial(self):
        cs = extract_constraints("TFT contrast ratio ≥400:1")
        ms = extract_measurements("TFT contrast 500 @25°C  410 @70°C  380 @85°C")
        vs = compare_constraint(cs[0], ms)
        assert summarize_verdicts(vs) == "PARTIELLEMENT_CONFORME"
        statuses = { (v.condition, v.status) for v in vs }
        assert any(v.status == "CONFORME" and "25" in v.condition for v in vs)
        assert any(v.status == "NON_CONFORME" and "85" in v.condition for v in vs)

    def test_cpu_load_conforme(self):
        cs = extract_constraints("CPU load <70%")
        ms = extract_measurements("CPU estimation = 50%")
        vs = compare_constraint(cs[0], ms)
        assert summarize_verdicts(vs) == "CONFORME"

    def test_temperature_fails_without_nok_word(self):
        cs = extract_constraints("Display temperature <38°C")
        ms = extract_measurements("Display surface 41.6°C")
        vs = compare_constraint(cs[0], ms)
        assert summarize_verdicts(vs) == "NON_CONFORME"


class TestContradictionRules:

    def test_tdr_polarity(self):
        assert tdr_polarity("Typ 119.9 mA NOK, Deviation") == "FAIL"
        assert tdr_polarity("CPU estimation = 50% meets the target") == "PASS"
        assert tdr_polarity("no polarity here 50 mA") == "NONE"

    def test_claimed_ok_no_evidence_is_an_assertion(self):
        c = classify_contradiction(
            req_id="REQ-1", matrix_status="OK", evidence_status="MANQUANT",
            coherence="UNVERIFIABLE", target="<=0.3 mm", supplier_result="",
            comment="OK", evidence_excerpt="", evidence_location="",
            domain="Mechanical", confidence="NONE",
            constraints=[], measurements=[],
        )
        assert c and c.type == "CLAIM_OK_NO_EVIDENCE"

    def test_claimed_nok_but_tdr_passes(self):
        c = classify_contradiction(
            req_id="REQ-2", matrix_status="NOK", evidence_status="CONFORME",
            coherence="MATRIX_TOO_PESSIMISTIC", target="<70 %",
            supplier_result="50 %", comment="NOK", evidence_excerpt="CPU 50%",
            evidence_location="Slide 30", domain="Software / CPU",
            confidence="HIGH", constraints=[], measurements=[],
        )
        assert c and c.type == "CLAIM_NOK_EVIDENCE_PASSES"


class TestStatementCrosswalk:

    def test_restated_target_extracted(self):
        ts = extract_restated_targets("Startup time target 800ms, measured 450ms")
        assert ts
        assert any(abs(t.value - 800) < 0.1 and t.unit_family == "time" for t in ts)

    def test_measured_skips_restated_target(self):
        ms = extract_measurements("Startup time target 800ms measured 450ms")
        times = [m for m in ms if m.unit_family == "time"]
        assert any(abs(m.value - 450) < 0.1 for m in times)
        assert not any(abs(m.value - 800) < 0.1 for m in times)

    def test_negative_temperature_constraint(self):
        cs = extract_constraints("Storage temperature ≥-40°C")
        assert cs
        assert cs[0].unit_family == "temperature"
        assert cs[0].value == -40
        assert cs[0].operator == "ge"

    def test_matrix_comment_vs_tdr_value_mismatch(self):
        comment_m = extract_measurements("Typ 40 mA — OK")
        tdr_m = extract_measurements("Idle current Typ 80 mA NOK")
        assert values_relation(comment_m, tdr_m) == "DISAGREE"

    def test_matrix_polarity_ok_vs_comment_nok(self):
        assert matrix_polarity("OK", "Typ 40 mA — OK") == "PASS"
        assert matrix_polarity("NOK", "Deviation") == "FAIL"
        assert matrix_polarity("OK", "NOK Deviation") == "FAIL"


class TestDeviationClassification:

    def test_plain_deviation(self):
        assert classify_conformity("DEVIATION") == "DEVIATION"
        assert classify_conformity("Deviation accepted") == "DEVIATION"

    def test_no_deviation_is_not_deviation(self):
        assert classify_conformity("no deviation") != "DEVIATION"

    def test_ok_and_nok_unchanged(self):
        assert classify_conformity("OK") == "OK"
        assert classify_conformity("NOK") == "NOK"
        assert classify_conformity("EE: ok") == "OK"


# ── End-to-end pipeline on synthetic Tianma-style files ───────────

TDR_TEXT = """\
Slide 12: EE Power Consumption
Reduced Consumption Mode current
Typ 119.9 mA
Max 192.7 mA
NOK, Deviation

Slide 18: LCF optical performance
lum attenuates >85% @V=32°
lum attenuates >95% @V=44°
NOK versus the 95% @32° target

Slide 24: TFT contrast ratio
500 @25°C
410 @70°C
380 @85°C

Slide 30: Software CPU
CPU estimation = 50%

Slide 41: Thermal
Display temperature 41.6°C
Deviation

Slide 50: Idle current
Idle current Typ 80 mA
NOK

Slide 55: EMC immunity
EMC immunity test NOK Deviation

Slide 62: Startup time
Startup time target 800ms
measured 450ms
OK

Slide 64: Standby current
Standby current 18 mA OK

Slide 65: Standby current second test
Standby current 35 mA

Slide 72: Storage temperature
Storage temperature tested -40°C OK
"""


def _build_matrix(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Matrix"
    ws.append([
        "Numéro de l'exigence", "Référence", "", "", "",
        "Libellé", "Conformité FNR", "Commentaires FNR",
    ])
    rows = [
        ("REQ-0308287", "EE",
         "Current consumption in Reduced Consumption Mode ≤100mA",
         "NOK", "Typ 119.9 mA / Max 192.7 mA — Deviation"),
        ("REQ-0307942", "OPT",
         "LCF Attenuation >95% at V=32°",
         "NOK", ">85% @V=32°, >95% only @V=44°"),
        ("REQ-0308444", "OPT",
         "TFT contrast ratio ≥400:1",
         "OK", "500@25°C / 410@70°C / 380@85°C"),
        ("REQ-0309001", "SW",
         "CPU load <70%",
         "OK", "estimated 50%"),
        ("REQ-0309002", "ME",
         "Display temperature <38°C",
         "OK", "41.6°C measured"),
        ("REQ-0309003", "ME",
         "Mechanical gap shall be ≤0.3 mm",
         "OK", "not shown in this TDR extract"),
        ("REQ-0309004", "SYS",
         "Diagnostic session is not applicable on this variant",
         "NA", "NA"),
        ("REQ-0309100", "EE",
         "Idle current ≤50mA",
         "OK", "Typ 40 mA — OK"),
        ("REQ-0309200", "EMC",
         "EMC immunity shall be compliant",
         "OK", "OK"),
        ("REQ-0309300", "EE",
         "Standby current ≤20mA",
         "OK", "18 mA"),
        ("REQ-0309400", "SYS",
         "Startup time ≤500ms",
         "OK", "meets spec"),
        ("REQ-0309500", "ME",
         "Storage temperature ≥-40°C",
         "DEVIATION", "we only guarantee -30°C"),
    ]
    for req, ref, desc, status, comment in rows:
        ws.append([req, ref, "", "", "", desc, status, comment])
    wb.save(path)
    return path


@pytest.fixture
def tianma_pair(tmp_path):
    matrix = _build_matrix(tmp_path / "tianma_matrix.xlsx")
    tdr = tmp_path / "S_12F_Tianma_TDR_report.txt"
    tdr.write_text(TDR_TEXT, encoding="utf-8")
    return matrix, tdr


class TestAerisPipeline:

    def test_matrix_parser_reads_fixture(self, tianma_pair):
        matrix, _ = tianma_pair
        analysis = extract_conformity_data(str(matrix), matrix.name)
        ids = {i.req_id for i in analysis.items}
        assert "REQ-0308287" in ids
        assert analysis.stats.get("NOK", 0) >= 2

    def test_txt_tdr_keeps_slide_locations(self):
        doc = parse_evidence_bytes("tdr.txt", TDR_TEXT.encode("utf-8"))
        locs = {c.location for c in doc.chunks}
        assert "Slide 12" in locs
        assert "Slide 24" in locs

    def test_three_conversation_examples(self, tianma_pair):
        matrix, tdr = tianma_pair
        report = run_crosscheck(str(matrix), [(tdr.name, tdr.read_bytes())], matrix.name)
        by_id = {i.req_id: i for i in report.items}

        current = by_id["REQ-0308287"]
        assert current.final_status in ("NON_CONFORME", "DEVIATION")
        assert current.confidence in ("VERY_HIGH", "HIGH")
        assert "12" in current.evidence_location or "current" in current.evidence_excerpt.lower()
        assert current.coherence == "ALIGNED"  # matrix already NOK

        lcf = by_id["REQ-0307942"]
        assert lcf.final_status == "PREUVE_INSUFFISANTE"
        assert not lcf.gap
        assert "85" in lcf.supplier_result

        contrast = by_id["REQ-0308444"]
        assert contrast.final_status == "PARTIELLEMENT_CONFORME"
        assert contrast.coherence == "MATRIX_TOO_OPTIMISTIC"
        conds = {v["status"] for v in contrast.condition_verdicts}
        assert "CONFORME" in conds and "NON_CONFORME" in conds

    def test_cpu_conforme_and_thermal_miss_without_nok_word(self, tianma_pair):
        matrix, tdr = tianma_pair
        report = run_crosscheck(str(matrix), [(tdr.name, tdr.read_bytes())], matrix.name)
        by_id = {i.req_id: i for i in report.items}

        assert by_id["REQ-0309001"].final_status == "CONFORME"
        thermal = by_id["REQ-0309002"]
        assert thermal.final_status == "NON_CONFORME"
        assert thermal.coherence == "MATRIX_TOO_OPTIMISTIC"

    def test_missing_evidence_is_manquant_not_invented(self, tianma_pair):
        matrix, tdr = tianma_pair
        report = run_crosscheck(str(matrix), [(tdr.name, tdr.read_bytes())], matrix.name)
        mech = next(i for i in report.items if i.req_id == "REQ-0309003")
        assert mech.final_status in ("MANQUANT", "PREUVE_INSUFFISANTE")
        assert mech.coherence == "UNVERIFIABLE"

    def test_contradiction_queue_flags_claimed_ok_vs_tdr(self, tianma_pair):
        matrix, tdr = tianma_pair
        report = run_crosscheck(str(matrix), [(tdr.name, tdr.read_bytes())], matrix.name)
        types = {}
        for c in report.contradictions:
            types.setdefault(c.req_id, set()).add(c.type)

        assert "CLAIM_OK_EVIDENCE_FAILS" in types["REQ-0309002"]
        assert "CLAIM_OK_PARTIAL" in types["REQ-0308444"]
        assert "CLAIM_OK_EVIDENCE_FAILS" in types["REQ-0309100"]
        assert "VALUE_MISMATCH" in types["REQ-0309100"]
        assert "CLAIM_OK_NO_EVIDENCE" in types["REQ-0309003"]
        assert "CLAIM_OK_TDR_SAYS_NOK" in types["REQ-0309200"]

        # Honest NOK in the matrix that the TDR confirms is NOT a contradiction.
        assert "REQ-0308287" not in types
        assert "REQ-0307942" not in types

    def test_statement_crosswalk_finds_said_incompliances(self, tianma_pair):
        matrix, tdr = tianma_pair
        report = run_crosscheck(str(matrix), [(tdr.name, tdr.read_bytes())], matrix.name)
        by_id = {w.req_id: w for w in report.crosswalk}
        types = {}
        for c in report.contradictions:
            types.setdefault(c.req_id, set()).add(c.type)

        # Honest NOK: same numbers, same polarity — not an incompliance.
        current = by_id["REQ-0308287"]
        assert current.alignment == "ALIGNED"
        assert current.values_agree == "AGREE"

        idle = by_id["REQ-0309100"]
        assert idle.values_agree == "DISAGREE"
        assert idle.alignment == "VALUE_MISMATCH"

        startup = by_id["REQ-0309400"]
        assert startup.restated_target_match == "WRONG_TARGET"
        assert "TDR_RESTATES_WRONG_TARGET" in types["REQ-0309400"]
        # Measured 450ms still meets ≤500ms — the lie is the restated limit.
        startup_item = next(i for i in report.items if i.req_id == "REQ-0309400")
        assert startup_item.final_status == "CONFORME"

        standby_types = types["REQ-0309300"]
        assert "TDR_INTERNAL_CONFLICT" in standby_types
        assert "CLAIM_OK_PARTIAL" in standby_types

        storage = next(i for i in report.items if i.req_id == "REQ-0309500")
        assert storage.final_status == "CONFORME"
        assert storage.coherence == "MATRIX_TOO_PESSIMISTIC"

        assert report.crosswalk_summary["valueMismatch"] >= 1
        assert report.crosswalk_summary["wrongTarget"] >= 1
        assert report.crosswalk_summary["tdrConflict"] >= 1
        assert report.deviations
        assert report.tdr_statements
        assert report.coverage.get("total") == 12

    def test_incoherences_are_one_readable_row_per_requirement(self, tianma_pair):
        matrix, tdr = tianma_pair
        report = run_crosscheck(str(matrix), [(tdr.name, tdr.read_bytes())], matrix.name)
        rows = report.incoherences

        # Une seule ligne par exigence, même si plusieurs motifs.
        ids = [r["req_id"] for r in rows]
        assert len(ids) == len(set(ids))

        # Les exigences cohérentes ne sont jamais listées.
        assert "REQ-0309001" not in ids   # CPU 50% < 70%, matrice OK
        assert "REQ-0309004" not in ids   # NA
        assert "REQ-0308287" not in ids   # NOK assumé, TDR confirme

        # Tri par gravité : les bloquantes d'abord.
        assert rows[0]["gravite"].startswith("1 - Bloquant")
        assert [r["n"] for r in rows] == list(range(1, len(rows) + 1))

        thermal = next(r for r in rows if r["req_id"] == "REQ-0309002")
        assert thermal["demande"] == "<38 °C"
        assert "41.6" in thermal["tdr"]
        assert thermal["ou"] == "Slide 41"
        assert "déclare OK" in thermal["pourquoi"]
        assert "CLAIM_OK" not in thermal["pourquoi"]
        assert thermal["action"]

        idle = next(r for r in rows if r["req_id"] == "REQ-0309100")
        assert idle["demande"] == "≤50 mA"
        assert "80 mA" in idle["tdr"]
        assert idle["autres_motifs"]  # value mismatch listé en motif secondaire

        # Le signe négatif ne doit jamais être perdu.
        storage = next(r for r in rows if r["req_id"] == "REQ-0309500")
        assert "-40 °C" in storage["tdr"] or "-30" in storage["pourquoi"]

        assert report.incoherence_summary["total"] == len(rows)
        assert report.incoherence_summary["bloquant"] >= 1

    def test_na_stays_na(self, tianma_pair):
        matrix, tdr = tianma_pair
        report = run_crosscheck(str(matrix), [(tdr.name, tdr.read_bytes())], matrix.name)
        na = next(i for i in report.items if i.req_id == "REQ-0309004")
        assert na.final_status == "NA"

    def test_synthesis_counts_and_excel(self, tianma_pair):
        matrix, tdr = tianma_pair
        report = run_crosscheck(str(matrix), [(tdr.name, tdr.read_bytes())], matrix.name)
        payload = report_to_dict(report)
        s = payload["summary"]
        assert s["total"] == 12
        assert s["nonConforme"] + s["deviation"] >= 3
        assert s["conforme"] >= 1
        assert s["partiel"] >= 1
        assert s["matrixTooOptimistic"] >= 2
        assert s["contradictions"] >= 3
        assert s["contradictionsCritical"] >= 1
        assert s["valueMismatch"] >= 1
        assert s["wrongTarget"] >= 1
        assert payload["topRisks"]
        assert payload["contradictions"]
        assert payload["crosswalk"]
        assert payload["tdrStatements"]
        assert payload["deviations"]
        assert payload["incoherences"]
        assert s["incoherences"] == len(payload["incoherences"])
        assert s["bloquant"] >= 1

        xlsx = generate_aeris_excel(payload)
        wb = load_workbook(io.BytesIO(xlsx))
        # Le classeur s'ouvre sur les incohérences, et rien d'autre n'est requis.
        assert wb.sheetnames == ["Incohérences", "Synthèse", "Détail complet"]
        assert wb.active.title == "Incohérences"
        ws = wb["Incohérences"]
        assert ws.cell(5, 1).value == "N°"
        assert ws.cell(5, 9).value == "Pourquoi c'est une incohérence"
        # Première ligne de données : une phrase, pas un code machine.
        pourquoi = ws.cell(6, 9).value
        assert pourquoi and "matrice" in pourquoi.lower()
        assert "CLAIM_OK" not in pourquoi

    def test_pptx_fallback_xml_roundtrip(self, tmp_path, tianma_pair):
        """A minimal PPTX (zip/XML) is readable even without python-pptx shapes."""
        pptx = _minimal_pptx_with_text(
            tmp_path / "tdr.pptx",
            "EE Power Consumption\nReduced mode current Typ 119.9mA Max 192.7mA",
        )
        matrix, _ = tianma_pair
        report = run_crosscheck(str(matrix), [(pptx.name, pptx.read_bytes())], matrix.name)
        current = next(i for i in report.items if i.req_id == "REQ-0308287")
        assert current.final_status in ("NON_CONFORME", "DEVIATION")
        assert current.evidence_location.startswith("Slide")


def _minimal_pptx_with_text(path: Path, text: str) -> Path:
    """Write a tiny valid PPTX so the zip/XML extractor has something to read."""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(4))
        box.text_frame.text = text
        prs.save(path)
        return path
    except Exception:
        pytest.skip("python-pptx not installed — XML fallback covered by TXT slides")


# ── HTTP route (optional — needs FastAPI test client) ─────────────

class TestAerisRoute:
    def test_endpoint_returns_synthesis(self, tianma_pair):
        try:
            from fastapi.testclient import TestClient
            from app.conformity_server import app
        except Exception:
            pytest.skip("FastAPI TestClient not available")

        matrix, tdr = tianma_pair
        client = TestClient(app)
        with matrix.open("rb") as mf, tdr.open("rb") as ef:
            res = client.post(
                "/api/aeris-crosscheck",
                files=[
                    ("matrix", (matrix.name, mf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
                    ("evidence", (tdr.name, ef, "text/plain")),
                ],
            )
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["summary"]["total"] == 12
        assert body["reportExcel"]
        ids = {i["req_id"] for i in body["items"]}
        assert "REQ-0308287" in ids
