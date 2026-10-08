"""Request-scoped version comparison route for the public AERIS UI."""
from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from fastapi import APIRouter, File, HTTPException, UploadFile

router = APIRouter(prefix="/api", tags=["conformity-comparison"])
MAX_UPLOAD_BYTES = 40 * 1024 * 1024
MATRIX_EXTENSIONS = {".ods", ".xlsx", ".xlsm"}


def _excel_value(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    text = str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _fallback_delta_report(result: dict, names: list[str]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Delta Overview"
    overview.append(["AERIS — Version Comparison"])
    overview.append(["Versions", _excel_value(" → ".join(names))])
    overview.append(["Requirements compared", _excel_value(result.get("totalCompared", ""))])
    overview.append(["Status changes", _excel_value(result.get("totalChanges", ""))])
    overview.append(["Requirements missing", _excel_value(result.get("totalMissing", ""))])

    changes = workbook.create_sheet("Status Changes")
    headers = ["Step", "Requirement", "Reference", "From", "To", "Change"]
    changes.append(headers)
    for item in result.get("statusChanges") or []:
        if isinstance(item, dict):
            changes.append(
                [
                    _excel_value(item.get("step", "")),
                    _excel_value(item.get("reqId", "")),
                    _excel_value(item.get("reference", "")),
                    _excel_value(item.get("from", "")),
                    _excel_value(item.get("to", "")),
                    _excel_value(item.get("changeType", "")),
                ]
            )

    report = workbook.create_sheet("Review Notes")
    report.append(["Comparison summary"])
    report.append([_excel_value(result.get("reportText", "Comparison completed."))])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.post("/aeris-conformity-compare")
async def compare_conformity_versions(files: list[UploadFile] = File(...)) -> dict:
    if len(files) < 2:
        raise HTTPException(status_code=400, detail="Upload at least two matrix versions")

    names: list[str] = []
    paths: list[str] = []
    with TemporaryDirectory(prefix="aeris-compare-") as temporary_directory:
        for index, file in enumerate(files):
            name = (file.filename or "").replace("\\", "/").rsplit("/", 1)[-1]
            if not name or Path(name).suffix.lower() not in MATRIX_EXTENSIONS:
                accepted = ", ".join(sorted(MATRIX_EXTENSIONS))
                raise HTTPException(
                    status_code=400,
                    detail=f"Unsupported matrix file. Accepted: {accepted}",
                )

            content = await file.read(MAX_UPLOAD_BYTES + 1)
            if len(content) > MAX_UPLOAD_BYTES:
                raise HTTPException(status_code=413, detail="Matrix file exceeds 40 MB")
            if not content:
                raise HTTPException(status_code=400, detail=f"File '{name}' is empty")

            path = Path(temporary_directory) / f"{index}_{name}"
            path.write_bytes(content)
            names.append(name)
            paths.append(str(path))

        try:
            from app.qa.conformity_analyzer import compare_matrices, comparison_to_dict
            from app.qa.conformity_report import generate_delta_excel

            comparison = compare_matrices(paths, names)
            result = comparison_to_dict(comparison)
            steps = getattr(comparison, "steps", None)
            report = (
                generate_delta_excel(comparison)
                if steps is not None
                else _fallback_delta_report(result, names)
            )
            result.setdefault("steps", [])
            result["reportExcel"] = base64.b64encode(
                report
            ).decode("ascii")
            return result
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Comparison failed: {exc}") from exc
