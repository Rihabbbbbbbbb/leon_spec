"""Excel / Word exports of a multi-supplier TDR benchmark result."""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

T = {
    "fr": {
        "synthesis": "Synthèse", "scores": "Scores", "domains": "Comparaison domaines",
        "params": "Paramètres clés", "profiles": "Profils fournisseurs", "questions": "Questions",
        "facts": "Faits sources", "docs": "Documents", "supplier": "Fournisseur", "rank": "Rang",
        "score": "Score pondéré /5", "percent": "Score %", "verdict": "Verdict", "headline": "En bref",
        "domain": "Domaine", "weight": "Poids", "summary": "Résumé", "differentiators": "Différenciateurs",
        "assessment": "Évaluation", "coverage": "Couverture", "strengths": "Forces", "weaknesses": "Faiblesses",
        "risks": "Risques", "parameter": "Paramètre", "question": "Question", "priority": "Priorité",
        "rationale": "Justification", "page": "Page", "kind": "Type", "topic": "Sujet", "value": "Valeur",
        "variant": "Variante", "statement": "Énoncé", "quote": "Citation source", "importance": "Importance",
        "grounding": "Vérification", "file": "Fichier", "pages": "Pages", "facts_n": "Faits",
        "vision": "Pages vision", "warnings": "Avertissements", "title": "Benchmark technique TDR multi-fournisseurs",
        "exec": "Synthèse exécutive", "reco": "Recommandation technique", "preferred": "Fournisseur préféré",
        "runners": "Suivants", "conditions": "Conditions / clarifications", "cross": "Constats transverses",
        "major_risks": "Risques majeurs", "next": "Prochaines étapes", "confidence": "Limites de l'analyse",
        "overview": "Vue d'ensemble", "positioning": "Positionnement", "deviations": "Écarts déclarés",
        "assumptions": "Hypothèses / dépendances", "open_points": "Points ouverts", "generated": "Généré le",
        "mode": "Mode", "differs": "Attention : la recommandation diffère du meilleur score pondéré",
        "aspect": "Point de comparaison", "best": "Meilleur(s)", "deviations_n": "Écarts", "open_n": "Points ouverts",
        "source": "Source", "facts_verified": "Faits vérifiés",
        "overrides": "Scores / poids ajustés par l'expert : les totaux et le classement sont recalculés "
                     "(les textes IA reflètent les scores initiaux).",
        "comments": "Commentaires expert",
        "disclaimer": "Aide à la décision technique générée par IA à partir des seuls documents fournis "
                      "(offres techniques, hors aspects commerciaux). Chaque constat renvoie à des faits sourcés "
                      "(page + citation vérifiée) : à valider par les experts.",
    },
    "en": {
        "synthesis": "Synthesis", "scores": "Scores", "domains": "Domain comparison", "params": "Key parameters",
        "profiles": "Supplier profiles", "questions": "Questions", "facts": "Source facts", "docs": "Documents",
        "supplier": "Supplier", "rank": "Rank", "score": "Weighted score /5", "percent": "Score %",
        "verdict": "Verdict", "headline": "Headline", "domain": "Domain", "weight": "Weight", "summary": "Summary",
        "differentiators": "Differentiators", "assessment": "Assessment", "coverage": "Coverage",
        "strengths": "Strengths", "weaknesses": "Weaknesses", "risks": "Risks", "parameter": "Parameter",
        "question": "Question", "priority": "Priority", "rationale": "Rationale", "page": "Page", "kind": "Kind",
        "topic": "Topic", "value": "Value", "variant": "Variant", "statement": "Statement", "quote": "Source quote",
        "importance": "Importance", "grounding": "Verification", "file": "File", "pages": "Pages",
        "facts_n": "Facts", "vision": "Vision pages", "warnings": "Warnings",
        "title": "Multi-supplier TDR technical benchmark", "exec": "Executive summary",
        "reco": "Technical recommendation", "preferred": "Preferred supplier", "runners": "Runners-up",
        "conditions": "Conditions / clarifications", "cross": "Cross-cutting findings", "major_risks": "Major risks",
        "next": "Next steps", "confidence": "Limits of the analysis", "overview": "Overview",
        "positioning": "Positioning", "deviations": "Declared deviations", "assumptions": "Assumptions / dependencies",
        "open_points": "Open points", "generated": "Generated", "mode": "Mode",
        "differs": "Warning: the recommendation differs from the best weighted score",
        "aspect": "Comparison point", "best": "Best", "deviations_n": "Deviations", "open_n": "Open points",
        "source": "Source", "facts_verified": "Verified facts",
        "overrides": "Scores / weights adjusted by the expert: totals and ranking are recomputed "
                     "(AI texts reflect the initial scores).",
        "comments": "Expert comments",
        "disclaimer": "AI-generated technical decision support based only on the supplied documents (technical "
                      "offers, commercial aspects excluded). Every finding links to sourced facts (page + verified "
                      "quote): to be validated by experts.",
    },
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=14, bold=True, color="1F3864")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _t(result: dict) -> dict:
    return T["en" if result.get("language") == "en" else "fr"]


def _names(result: dict) -> dict[str, str]:
    return {s["id"]: s["name"] for s in result["suppliers"]}


def _cited(items: list[dict], key: str = "text") -> str:
    lines = []
    for item in items or []:
        refs = ", ".join(item.get("fact_ids") or [])
        sev = f"[{item['severity']}] " if item.get("severity") else ""
        lines.append(f"• {sev}{item.get(key, '')}" + (f" ({refs})" if refs else ""))
    return "\n".join(lines)


def _header(ws, row: int, values: list[str]) -> None:
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER


def _row(ws, row: int, values: list[Any]) -> None:
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.alignment = WRAP
        cell.border = BORDER


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_excel(result: dict) -> bytes:
    t = _t(result)
    names = _names(result)
    sids = [s["id"] for s in result["suppliers"]]
    ranked = sorted(result["suppliers"], key=lambda s: s["rank"])
    syn = result["synthesis"]
    wb = Workbook()

    # Synthesis
    ws = wb.active
    ws.title = t["synthesis"][:31]
    ws["A1"] = f"{t['title']} – {result.get('title', '')}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"{t['generated']}: {result.get('generatedAt', '')} | {t['mode']}: {result.get('mode')}"
    ws["A3"] = t["disclaimer"]
    ws["A3"].alignment = WRAP
    ws.merge_cells("A3:F3")
    ws.row_dimensions[3].height = 45
    if result.get("overridesApplied"):
        ws["A4"] = t["overrides"]
        ws["A4"].font = Font(bold=True, color="C00000")
    _header(ws, 5, [t["rank"], t["supplier"], t["score"], t["percent"], t["verdict"], t["headline"]])
    r = 6
    for s in ranked:
        v = syn["verdicts"].get(s["id"], {})
        _row(ws, r, [s["rank"], s["name"], s["score"], s["percent"],
                     result["verdictLabels"].get(v.get("verdict"), v.get("verdict", "")), v.get("headline", "")])
        r += 1
    r += 1
    rec = syn["recommendation"]
    blocks = [
        (t["exec"], syn.get("executiveSummary", "")),
        (t["reco"], f"{t['preferred']}: {names.get(rec.get('preferred'), '-')}\n"
                    f"{t['runners']}: {', '.join(names.get(x, x) for x in rec.get('runnersUp', []))}\n"
                    f"{rec.get('rationale', '')}" + (f"\n⚠ {t['differs']}" if rec.get("differsFromScore") else "")),
        (t["conditions"], "\n".join(f"• {c}" for c in rec.get("conditions", []))),
        (t["cross"], "\n".join(f"• {c}" for c in syn.get("crossCutting", []))),
        (t["major_risks"], "\n".join(f"• [{x.get('severity')}] {names.get(x.get('supplier_id'), '')}: {x.get('text')}"
                                     for x in syn.get("majorRisks", []))),
        (t["next"], "\n".join(f"• {c}" for c in syn.get("nextSteps", []))),
        (t["confidence"], syn.get("confidenceNote", "")),
    ]
    comments = (result.get("overrides") or {}).get("comments") or {}
    if comments:
        blocks.append((t["comments"], "\n".join(f"• {names.get(k, k)}: {v}" for k, v in comments.items())))
    for label, text in blocks:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=text)
        cell.alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = max(30, min(400, 15 * (len(text) // 110 + text.count("\n") + 1)))
        r += 1
    _widths(ws, [24, 26, 16, 10, 26, 70])

    # Scores heatmap
    ws = wb.create_sheet(t["scores"][:31])
    _header(ws, 1, [t["domain"], t["weight"]] + [names[s] for s in sids])
    r = 2
    for d in result["domains"]:
        _row(ws, r, [d["label"], d["weight"]] + [result["scoreMatrix"][d["key"]][s] for s in sids])
        r += 1
    _row(ws, r, [t["score"], ""] + [result["totals"][s]["score"] for s in sids])
    for c in range(1, len(sids) + 3):
        ws.cell(row=r, column=c).font = Font(bold=True)
    last_col = get_column_letter(len(sids) + 2)
    ws.conditional_formatting.add(f"C2:{last_col}{r}", ColorScaleRule(
        start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=2.5,
        mid_color="FFEB84", end_type="num", end_value=5, end_color="63BE7B"))
    r += 2
    _header(ws, r, [t["domain"], ""] + [f"{names[s]} – {t['facts_n']}/{t['deviations_n']}/{t['open_n']}" for s in sids])
    r += 1
    for d in result["domains"]:
        cov = result["coverage"][d["key"]]
        _row(ws, r, [d["label"], ""] + [f"{cov[s]['facts']} / {cov[s]['deviations']} / {cov[s]['openPoints']}"
                                         for s in sids])
        r += 1
    _widths(ws, [36, 8] + [20] * len(sids))
    ws.freeze_panes = "C2"

    # Domain comparison
    ws = wb.create_sheet(t["domains"][:31])
    _header(ws, 1, [t["domain"], t["supplier"], "Score", t["coverage"], t["assessment"], t["strengths"],
                    t["weaknesses"], t["risks"]])
    r = 2
    for d in result["domains"]:
        dr = result["domainResults"][d["key"]]
        _row(ws, r, [d["label"], t["summary"], "", "", dr.get("summary", ""),
                     "\n".join(f"• {x}" for x in dr.get("differentiators", [])), "", ""])
        ws.cell(row=r, column=1).font = Font(bold=True)
        r += 1
        for p in dr.get("points", []):
            text = "\n".join(f"• {names.get(x['supplier_id'], x['supplier_id'])}: {x['position']}"
                             + (f" ({', '.join(x['fact_ids'])})" if x["fact_ids"] else "") for x in p["positions"])
            best = ", ".join(names.get(x, x) for x in p.get("best_supplier_ids", []))
            _row(ws, r, ["", f"{t['aspect']}: {p['aspect']}", "", best, text, "", "", ""])
            r += 1
        for sid in dr["ranking"]:
            a = dr["assessments"][sid]
            _row(ws, r, ["", names[sid], a["score"], a["coverage"], a.get("summary", ""),
                         _cited(a.get("strengths")), _cited(a.get("weaknesses")), _cited(a.get("risks"))])
            r += 1
    _widths(ws, [28, 26, 7, 12, 60, 50, 50, 50])
    ws.freeze_panes = "B2"

    # Parameters
    ws = wb.create_sheet(t["params"][:31])
    _header(ws, 1, [t["parameter"]] + [names[s] for s in sids])
    r = 2
    for key, cells in result["parameters"].items():
        label = result["parameterLabels"][key]
        values = []
        for s in sids:
            entries = cells.get(s) or []
            values.append("\n".join(
                f"{e['value']}" + (f" [{e['variant']}]" if e.get("variant") else "") + f" (p.{e['page']}"
                + (", regex" if e["source"] == "pattern" else "") + ")" for e in entries))
        _row(ws, r, [f"{label['label']}" + (f" ({label['unit']})" if label["unit"] else "")] + values)
        r += 1
    _widths(ws, [30] + [38] * len(sids))
    ws.freeze_panes = "B2"

    # Profiles
    ws = wb.create_sheet(t["profiles"][:31])
    _header(ws, 1, [t["supplier"], t["overview"], t["positioning"], t["strengths"], t["weaknesses"], t["risks"],
                    t["deviations"], t["assumptions"], t["open_points"]])
    r = 2
    for s in ranked:
        p = result["profiles"][s["id"]]
        _row(ws, r, [s["name"], p.get("overview", ""), p.get("positioning", ""), _cited(p.get("strengths")),
                     _cited(p.get("weaknesses")), _cited(p.get("risks")), _cited(p.get("deviations")),
                     _cited(p.get("assumptions")), _cited(p.get("openPoints"))])
        r += 1
    _widths(ws, [18, 50, 40, 50, 50, 50, 50, 40, 40])

    # Questions
    ws = wb.create_sheet(t["questions"][:31])
    _header(ws, 1, [t["supplier"], t["domain"], t["priority"], t["question"], t["rationale"]])
    labels = {d["key"]: d["label"] for d in result["domains"]}
    r = 2
    for s in ranked:
        for q in result["profiles"][s["id"]].get("questions", []):
            _row(ws, r, [s["name"], labels.get(q.get("domain"), q.get("domain", "")), q.get("priority", ""),
                         q.get("question", ""), q.get("rationale", "")])
            r += 1
    _widths(ws, [18, 28, 10, 80, 50])
    ws.auto_filter.ref = f"A1:E{max(r - 1, 1)}"

    # Facts
    ws = wb.create_sheet(t["facts"][:31])
    docs = {d["id"]: d["fileName"] for d in result["documents"]}
    _header(ws, 1, ["ID", t["supplier"], t["file"], t["page"], t["domain"], t["kind"], t["importance"], t["topic"],
                    t["parameter"], t["value"], t["variant"], t["statement"], t["quote"], t["grounding"], t["source"]])
    r = 2
    for f in result["facts"]:
        _row(ws, r, [f["id"], names.get(f["supplier_id"]), docs.get(f["doc_id"]), f["page"],
                     labels.get(f["domain"], f["domain"]), result["factKinds"].get(f["kind"], f["kind"]),
                     f["importance"], f.get("topic", ""),
                     result["parameterLabels"].get(f.get("parameter"), {}).get("label", "") if f.get("parameter") else "",
                     f.get("value", ""), f.get("variant", ""), f["statement"], f["quote"], f["grounding"], f["source"]])
        r += 1
    _widths(ws, [11, 14, 30, 6, 22, 18, 9, 22, 18, 18, 12, 60, 60, 11, 8])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:O{max(r - 1, 1)}"

    # Documents
    ws = wb.create_sheet(t["docs"][:31])
    _header(ws, 1, [t["file"], t["supplier"], t["pages"], "Chars", t["vision"], t["facts_n"], t["warnings"]])
    r = 2
    for d in result["documents"]:
        n = sum(1 for f in result["facts"] if f["doc_id"] == d["id"])
        _row(ws, r, [d["fileName"], names.get(d["supplierId"]), d["pages"], d["chars"], d["visionPages"], n,
                     "\n".join(d.get("warnings") or [])])
        r += 1
    stats = result.get("stats", {})
    r += 1
    for k, v in stats.items():
        if not isinstance(v, dict):
            _row(ws, r, [k, v])
            r += 1
    for w in result.get("warnings", []):
        _row(ws, r, [t["warnings"], w])
        r += 1
    _widths(ws, [60, 18, 8, 10, 12, 8, 60])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_docx(result: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    t = _t(result)
    names = _names(result)
    sids = [s["id"] for s in result["suppliers"]]
    ranked = sorted(result["suppliers"], key=lambda s: s["rank"])
    syn = result["synthesis"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    doc.add_heading(t["title"], 0)
    if result.get("title"):
        doc.add_paragraph(result["title"])
    p = doc.add_paragraph(f"{t['generated']}: {result.get('generatedAt', '')} – {t['mode']}: {result.get('mode')}")
    p.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    note = doc.add_paragraph(t["disclaimer"])
    note.runs[0].italic = True
    if result.get("overridesApplied"):
        warn = doc.add_paragraph(t["overrides"])
        warn.runs[0].bold = True
        warn.runs[0].font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        for k, v in ((result.get("overrides") or {}).get("comments") or {}).items():
            doc.add_paragraph(f"{t['comments']} – {names.get(k, k)}: {v}", style="List Bullet")

    doc.add_heading(t["exec"], 1)
    doc.add_paragraph(syn.get("executiveSummary", ""))

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for i, h in enumerate([t["rank"], t["supplier"], t["score"], t["verdict"], t["headline"]]):
        table.rows[0].cells[i].text = h
    for s in ranked:
        v = syn["verdicts"].get(s["id"], {})
        cells = table.add_row().cells
        cells[0].text = str(s["rank"])
        cells[1].text = s["name"]
        cells[2].text = f"{s['score']} ({s['percent']}%)"
        cells[3].text = result["verdictLabels"].get(v.get("verdict"), v.get("verdict", "") or "")
        cells[4].text = v.get("headline", "") or ""

    rec = syn["recommendation"]
    doc.add_heading(t["reco"], 1)
    doc.add_paragraph(f"{t['preferred']}: {names.get(rec.get('preferred'), '-')}")
    if rec.get("runnersUp"):
        doc.add_paragraph(f"{t['runners']}: {', '.join(names.get(x, x) for x in rec['runnersUp'])}")
    doc.add_paragraph(rec.get("rationale", ""))
    if rec.get("differsFromScore"):
        doc.add_paragraph(f"⚠ {t['differs']}")
    for label, items in ((t["conditions"], rec.get("conditions", [])), (t["cross"], syn.get("crossCutting", [])),
                         (t["next"], syn.get("nextSteps", []))):
        if items:
            doc.add_heading(label, 2)
            for item in items:
                doc.add_paragraph(str(item), style="List Bullet")
    if syn.get("majorRisks"):
        doc.add_heading(t["major_risks"], 2)
        for x in syn["majorRisks"]:
            doc.add_paragraph(f"[{x.get('severity')}] {names.get(x.get('supplier_id'), '')}: {x.get('text')}",
                              style="List Bullet")

    doc.add_heading(t["scores"], 1)
    table = doc.add_table(rows=1, cols=len(sids) + 2)
    table.style = "Light Grid Accent 1"
    head = table.rows[0].cells
    head[0].text = t["domain"]
    head[1].text = t["weight"]
    for i, s in enumerate(sids):
        head[i + 2].text = names[s]
    for d in result["domains"]:
        cells = table.add_row().cells
        cells[0].text = d["label"]
        cells[1].text = str(d["weight"])
        for i, s in enumerate(sids):
            cells[i + 2].text = str(result["scoreMatrix"][d["key"]][s])
            cells[i + 2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cells = table.add_row().cells
    cells[0].text = t["score"]
    for i, s in enumerate(sids):
        cells[i + 2].text = str(result["totals"][s]["score"])

    doc.add_heading(t["domains"], 1)
    for d in result["domains"]:
        dr = result["domainResults"][d["key"]]
        if all(dr["assessments"][s]["score"] == 0 for s in sids):
            continue
        doc.add_heading(d["label"], 2)
        if dr.get("summary"):
            doc.add_paragraph(dr["summary"])
        for x in dr.get("differentiators", []):
            doc.add_paragraph(x, style="List Bullet")
        for sid in dr["ranking"]:
            a = dr["assessments"][sid]
            para = doc.add_paragraph()
            run = para.add_run(f"{names[sid]} – {a['score']}/5 ({a['coverage']}): ")
            run.bold = True
            para.add_run(a.get("summary", ""))
            for key in ("strengths", "weaknesses", "risks"):
                for item in a.get(key, [])[:4]:
                    refs = f" ({', '.join(item['fact_ids'])})" if item.get("fact_ids") else ""
                    doc.add_paragraph(f"{t[key]}: {item['text']}{refs}", style="List Bullet 2")

    doc.add_heading(t["profiles"], 1)
    for s in ranked:
        prof = result["profiles"][s["id"]]
        doc.add_heading(f"{s['name']} – {s['score']}/5", 2)
        if prof.get("overview"):
            doc.add_paragraph(prof["overview"])
        if prof.get("positioning"):
            doc.add_paragraph(prof["positioning"])
        for key, label in (("strengths", t["strengths"]), ("weaknesses", t["weaknesses"]), ("risks", t["risks"]),
                           ("deviations", t["deviations"]), ("openPoints", t["open_points"])):
            items = prof.get(key) or []
            if items:
                doc.add_paragraph().add_run(label).bold = True
                for item in items[:8]:
                    doc.add_paragraph(item.get("text", ""), style="List Bullet")
        questions = prof.get("questions") or []
        if questions:
            doc.add_paragraph().add_run(t["questions"]).bold = True
            for q in questions:
                doc.add_paragraph(f"[{q.get('priority', '')}] {q.get('question', '')}", style="List Number")

    if syn.get("confidenceNote"):
        doc.add_heading(t["confidence"], 1)
        doc.add_paragraph(syn["confidenceNote"])

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
