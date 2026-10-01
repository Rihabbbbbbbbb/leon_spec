"""Tests for the spec ↔ conformity matrix coverage & traceability report."""
import io
import sys
from pathlib import Path

import pytest

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

from app.qa.conformity_analyzer import ConformityAnalysis, ConformityItem, analyze_conformity_matrix
from app.qa.conformity_coverage import (
    build_coverage_report,
    canonical_key,
    coverage_to_dict,
    extract_id_tokens,
    normalize_id,
)
from app.qa.conformity_report import generate_coverage_excel

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "uploads"
ASU_SPEC = DATA_DIR / "00692_25_01250_ASU_Technical_Specification_SPX _1_.docx"
ASU_MATRIX = DATA_DIR / "00692_25_01250_ASU_Technical_Specification_SPX (1) (5)_Conformity_Matrix (2).xlsx"
DM12F = DATA_DIR / "2_01843_25_00540_PHYS_GEN_DM12F_Conformity_Matrix_v1_CarUX_DM12F_20260428_DONE 1.ods"

SPEC_TEXT = """
REF-ASU-CD-EXIFUNC-002 | The function shall manage the following.
REF-ASU-CD-EXIFUNC-003 | The function shall provide feedback.
GEN-HW-ST-SSC.009 | The hardware shall comply.
REQ-0937578 | The system shall detect.
"""


def _item(req_id="", reference="", category="OK", comment="", is_req=True, description=""):
    return ConformityItem(
        row_index=0, req_id=req_id, reference=reference,
        description=description, conformity_raw="", conformity_category=category,
        comment=comment, is_requirement=is_req,
    )


def _analysis(items):
    return ConformityAnalysis(items=items, file_name="matrix.xlsx")


class TestIdHelpers:
    def test_normalize_id(self):
        assert normalize_id("GEN-HW-ST-SSC.009(0)") == "GEN-HW-ST-SSC.009"
        assert normalize_id("ref-asu-cd-001") == "REF-ASU-CD-001"
        assert normalize_id("") == ""

    def test_extract_id_tokens(self):
        tokens = extract_id_tokens("[STA12] RA19 | [M8] GEN-HW-ST-SSC.009(0)")
        assert "GEN-HW-ST-SSC.009" in tokens

    def test_canonical_key_prefers_req_id(self):
        assert canonical_key("REQ-0307775", "[M8] GEN-HW-ST-SSC.009(0)") == "REQ-0307775"
        assert canonical_key("", "[M8] GEN-HW-ST-SSC.009(0)") == "GEN-HW-ST-SSC.009"
        assert canonical_key("", "") == ""


class TestCoverageReport:
    def test_matches_by_reference(self):
        analysis = _analysis([
            _item(req_id="REQ-1", reference="[M8] GEN-HW-ST-SSC.009(0)", category="OK", comment="ok"),
        ])
        report = build_coverage_report(SPEC_TEXT, analysis)
        assert report.spec_with_id == 4
        assert report.answered == 1
        assert report.matches[0]["matchedBy"] == "reference"
        assert report.matches[0]["specReqId"] == "GEN-HW-ST-SSC.009"
        assert report.missing == 3

    def test_matches_by_req_id(self):
        analysis = _analysis([
            _item(req_id="REQ-0937578", reference="", category="NOK", comment="nok"),
        ])
        report = build_coverage_report(SPEC_TEXT, analysis)
        assert report.answered == 1
        assert report.matches[0]["matchedBy"] == "req_id"
        assert report.matches[0]["category"] == "NOK"

    def test_matches_by_description_fallback(self):
        analysis = _analysis([
            _item(req_id="REQ-1", reference="", category="OK",
                  description="The function shall provide feedback to the driver within 200 milliseconds."),
        ])
        report = build_coverage_report(SPEC_TEXT, analysis)
        assert report.answered == 1
        assert report.matches[0]["matchedBy"] == "description"
        assert report.matches[0]["specReqId"] == "REF-ASU-CD-EXIFUNC-003"

    def test_unmatched_rows(self):
        analysis = _analysis([
            _item(req_id="REQ-9999999", reference="", category="OK", comment="x"),
        ])
        report = build_coverage_report(SPEC_TEXT, analysis)
        assert len(report.unmatched_rows) == 1
        assert report.unmatched_rows[0]["reqId"] == "REQ-9999999"

    def test_non_requirement_rows_ignored(self):
        analysis = _analysis([
            _item(req_id="", reference="Allocation matrix of the LVDS", category="OK", is_req=False),
        ])
        report = build_coverage_report(SPEC_TEXT, analysis)
        assert report.matrix_requirement_rows == 0
        assert report.answered == 0

    def test_coverage_rate(self):
        analysis = _analysis([
            _item(req_id="REQ-1", reference="[M8] GEN-HW-ST-SSC.009(0)", category="OK"),
            _item(req_id="REQ-2", reference="[M8] REF-ASU-CD-EXIFUNC-002(0)", category="OK"),
        ])
        report = build_coverage_report(SPEC_TEXT, analysis)
        assert report.answered == 2
        assert report.matched == 2
        assert report.coverage_rate == pytest.approx(0.5)
        assert report.answer_rate == pytest.approx(0.5)

    def test_unanswered_matrix_is_pending_not_answered(self):
        # The supplier returned the matrix with the requirement row present but
        # no verdict (EMPTY) — this must count as PENDING, not ANSWERED.
        analysis = _analysis([
            _item(req_id="REQ-1", reference="[M8] GEN-HW-ST-SSC.009(0)", category="EMPTY"),
        ])
        report = build_coverage_report(SPEC_TEXT, analysis)
        assert report.matched == 1
        assert report.answered == 0
        assert report.pending == 1
        assert report.answer_rate == 0.0
        assert report.coverage_rate == pytest.approx(0.25)
        assert report.matches[0]["category"] == "EMPTY"

    def test_to_dict_keys(self):
        analysis = _analysis([_item(req_id="REQ-1", reference="[M8] GEN-HW-ST-SSC.009(0)")])
        d = coverage_to_dict(build_coverage_report(SPEC_TEXT, analysis))
        for key in ("specTotal", "specWithId", "matchedCount", "answeredCount",
                    "pendingCount", "missingCount", "coverageRate", "answerRate",
                    "matches", "missing", "unmatchedRows", "reportText"):
            assert key in d

    def test_excel_generates(self):
        analysis = _analysis([
            _item(req_id="REQ-1", reference="[M8] GEN-HW-ST-SSC.009(0)", category="OK"),
            _item(req_id="REQ-2", reference="", category="NOK"),
        ])
        report = build_coverage_report(SPEC_TEXT, analysis)
        xlsx = generate_coverage_excel(report)
        assert xlsx[:2] == b"PK"
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(xlsx))
        assert "Coverage" in wb.sheetnames
        assert "Pending" in wb.sheetnames
        assert "Missing" in wb.sheetnames
        assert "Unmatched Rows" in wb.sheetnames


class TestAsuMatrix:
    """The ASU conformity matrix uses non-standard column names
    ('Engagement / Commitment', 'Commentaires / Comments', 'Statut PSA') and
    REF-ASU-… requirement ids — the analyzer must handle it and coverage must
    match the ASU spec by req_id."""

    @pytest.fixture(scope="class")
    def asu(self):
        if not (ASU_SPEC.exists() and ASU_MATRIX.exists()):
            pytest.skip("ASU spec/matrix not found")
        from app.qa.conformity_analyzer import extract_conformity_data
        from app.qa.retrieval import extract_text_from_file
        analysis = extract_conformity_data(str(ASU_MATRIX), ASU_MATRIX.name)
        spec_text = extract_text_from_file(ASU_SPEC)
        cov = build_coverage_report(spec_text, analysis,
                                    spec_name=ASU_SPEC.name, matrix_name=ASU_MATRIX.name)
        return analysis, cov

    def test_extracts_asu_matrix(self, asu):
        analysis, _ = asu
        assert len(analysis.items) > 0
        assert analysis.column_mapping["conformity"]  # 'Engagement' columns found
        assert analysis.column_mapping["comment"]     # 'Commentaires' columns found
        assert analysis.column_mapping["stellantis_verdict"]  # 'Statut PSA'
        # every row with a requirement id (REF-/APP-/GEN-) is a requirement
        req_rows = [it for it in analysis.items if it.req_id]
        assert req_rows
        assert all(it.is_requirement for it in req_rows)
        assert all(it.req_id.startswith(("REF-", "APP-", "GEN-")) for it in req_rows)

    def test_coverage_matches_asu_spec_by_req_id(self, asu):
        _, cov = asu
        # The ASU matrix is pre-filled from the same spec, so EVERY spec
        # requirement must be TRACED (matched) by req_id — the exact number of
        # answered rows varies as the supplier fills answers in over time, so
        # only the INVARIANTS are asserted, never a hardcoded count:
        #   answered + pending == matched == spec_with_id, missing == 0.
        assert cov.matched == cov.spec_with_id
        assert cov.missing == 0
        assert cov.answered + cov.pending == cov.matched
        assert cov.answered + cov.pending + cov.missing == cov.spec_with_id
        assert cov.coverage_rate == 1.0
        assert cov.answer_rate == pytest.approx(cov.answered / cov.spec_with_id, abs=1e-3)
        assert all(m["matchedBy"] == "req_id" for m in cov.matches)


class TestRealFiles:
    @pytest.fixture(scope="class")
    def real(self):
        if not (ASU_SPEC.exists() and DM12F.exists()):
            pytest.skip("Real spec/matrix not found")
        from app.qa.retrieval import extract_text_from_file
        spec_text = extract_text_from_file(ASU_SPEC)
        analysis = analyze_conformity_matrix(str(DM12F), DM12F.name)
        return build_coverage_report(spec_text, analysis,
                                     spec_name=ASU_SPEC.name, matrix_name=DM12F.name)

    def test_builds_on_real_files(self, real):
        assert real.spec_with_id > 0
        assert real.matrix_requirement_rows > 0
        assert 0.0 <= real.coverage_rate <= 1.0
        assert 0.0 <= real.answer_rate <= 1.0
        assert real.matched + real.missing == real.spec_with_id
        assert real.answered + real.pending == real.matched

    def test_matches_are_consistent(self, real):
        for m in real.matches:
            assert m["matchedBy"] in ("req_id", "reference", "description")
            assert m["category"] in ("OK", "NOK", "NA", "EMPTY")