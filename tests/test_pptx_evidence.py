from types import SimpleNamespace

from pptx import Presentation
from pptx.util import Inches

from app.qa.pptx_evidence import SlideEvidence, extract_pptx_evidence
from app.qa.pptx_evidence_matching import analyze_matrix_against_pptx, match_requirement_evidence


def test_pptx_extracts_slide_text_and_table_citations(tmp_path):
    path = tmp_path / "tdr.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    textbox.name = "Measured result"
    textbox.text_frame.text = "REQ-12345: measured input voltage is 12 V"
    table_shape = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(6), Inches(1))
    table_shape.table.cell(0, 0).text = "Test condition"
    table_shape.table.cell(0, 1).text = "25 C"
    presentation.save(path)

    evidence = extract_pptx_evidence(path)

    assert len(evidence) == 2
    assert evidence[0].slide_number == 1
    assert evidence[0].shape_name == "Measured result"
    assert "REQ-12345" in evidence[0].evidence_ids
    assert "25 C" in evidence[1].text


def test_pptx_retrieval_does_not_self_confirm_supplier_comment():
    requirement = {
        "reqId": "REQ-98765",
        "description": "The display contrast ratio shall be at least 400 to 1.",
        "comment": "Measured luminance uniformity was 95 percent at 43 degrees.",
    }
    evidence = [SlideEvidence("tdr.pptx", 3, "Results", "text",
                             "Measured luminance uniformity was 95 percent at 43 degrees.")]

    assert match_requirement_evidence(requirement, evidence) == []


def test_pptx_retrieval_does_not_treat_shared_numbers_as_support():
    requirement = {
        "reqId": "REQ-98765",
        "description": "The minimum contrast ratio shall be 400 to 1.",
    }
    evidence = [SlideEvidence(
        "tdr.pptx", 63, "White luminance and view cone",
        "text", "White luminance >500 cd/m2 at 25 degrees; >410 at 70 degrees; >380 at 85 degrees",
    )]

    assert match_requirement_evidence(requirement, evidence) == []


def test_pptx_retrieval_repairs_spacing_in_requirement_id():
    evidence = [SlideEvidence(
        "tdr.pptx", 4, "Test result", "text", "REQ - 0308287: current consumption 85 mA",
    )]

    matches = match_requirement_evidence(
        {"reqId": "REQ-0308287", "description": "Current consumption shall be below 100 mA."},
        evidence,
    )

    assert len(matches) == 1
    assert matches[0]["matchType"] == "requirement_id"


def test_pptx_analysis_tracks_matrix_row_and_supplier_status(tmp_path, monkeypatch):
    from app.qa import pptx_evidence_matching

    monkeypatch.setattr(pptx_evidence_matching, "extract_pptx_evidence", lambda _: [
        SlideEvidence("tdr.pptx", 2, "Results", "text", "REQ-1 result: 12 V", ["REQ-1"]),
    ])
    analysis = SimpleNamespace(file_name="matrix.xlsx", items=[SimpleNamespace(
        is_requirement=True, row_index=5, req_id="REQ-1", reference="",
        description="Input shall be 12 V.", conformity_category="OK",
        conformity_raw="OK", comment="",)])

    result = analyze_matrix_against_pptx(analysis, tmp_path / "tdr.pptx")

    assert result["requirements"][0]["matrixRowNumber"] == 6
    assert result["requirements"][0]["matrixStatus"] == "OK"
    assert result["summary"]["complianceVerdictsMade"] == 0
    assert result["decisionSummary"]["supplierDeclared"]["OK"] == 1