"""Tests for the version-to-version delta (multi-matrix comparison)."""
import io
import sys
from pathlib import Path

import pytest

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

from app.qa.conformity_analyzer import (
    ConformityItem,
    _canonical_key,
    _clean_comments,
    _compute_deltas,
    _is_bare_domain_code,
    compare_matrices,
    comparison_to_dict,
)
from app.qa.conformity_report import generate_delta_excel

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "uploads"
TIANMA_V1 = DATA_DIR / "01843_25_00540_PHYS_GEN_DM17F_Conformity_Matrix_v1_FR_EN_TIANMA_20260427.xlsm"
TIANMA_V2 = DATA_DIR / "01843_25_00540_PHYS_GEN_DM17F_Conformity_Matrix_v1_FR_EN_TIANMA_20260427 (1).xlsm"


def _row(req_id, reference, category, comment="", version=""):
    return {"reqId": req_id, "reference": reference, "category": category,
            "comment": comment, "version": version}


class TestBareDomainCodeFiltering:
    def test_bare_domain_codes_are_filtered(self):
        assert _is_bare_domain_code("SYS")
        assert _is_bare_domain_code("sw")
        assert _is_bare_domain_code("ME/EE")
        assert _is_bare_domain_code("VE, ME")
        assert _is_bare_domain_code("SYS | SW")
        assert _is_bare_domain_code("EE ME")

    def test_real_comments_are_kept(self):
        assert not _is_bare_domain_code("")
        assert not _is_bare_domain_code("EE\nSame with DM12.3\nTyp 119.9mA")
        assert not _is_bare_domain_code("Discussed in QIA. Follow the same requirement.")
        assert not _is_bare_domain_code("Not tested yet")

    def test_clean_comments_drops_domain_codes(self):
        cleaned = _clean_comments(["SYS", "SW", "", "Real comment here", "ME/EE"])
        assert cleaned == ["Real comment here"]


class TestCanonicalKey:
    def test_prefers_req_id(self):
        item = ConformityItem(row_index=0, req_id="REQ-1", reference="[M8] GEN-HW-ST-SSC.009(0)")
        assert _canonical_key(item) == "REQ-1"

    def test_falls_back_to_reference(self):
        item = ConformityItem(row_index=0, req_id="", reference="[M8] GEN-HW-ST-SSC.009(0)")
        assert _canonical_key(item) == "GEN-HW-ST-SSC.009"


class TestComputeDeltas:
    def test_status_changes_and_regressions(self):
        per = {
            "v1.xlsx": {
                "REQ-1": _row("REQ-1", "R1", "NOK"),
                "REQ-2": _row("REQ-2", "R2", "OK"),
                "REQ-3": _row("REQ-3", "R3", "OK"),
            },
            "v2.xlsx": {
                "REQ-1": _row("REQ-1", "R1", "OK"),   # improvement
                "REQ-2": _row("REQ-2", "R2", "NOK"),  # regression
                "REQ-3": _row("REQ-3", "R3", "OK"),
            },
        }
        sc, cc, vc, new_in, removed_in, steps = _compute_deltas(
            per, ["v1.xlsx", "v2.xlsx"], {"REQ-1", "REQ-2", "REQ-3"})
        assert len(sc) == 2
        by_id = {c["reqId"]: c for c in sc}
        assert by_id["REQ-1"]["improvement"] is True
        assert by_id["REQ-1"]["changeType"] == "improved"
        assert by_id["REQ-2"]["improvement"] is False
        assert by_id["REQ-2"]["changeType"] == "regressed"
        assert steps[0]["statusChanges"] == 2
        assert steps[0]["answeredBefore"] == 3
        assert steps[0]["answeredAfter"] == 3

    def test_empty_is_worst_and_answer_removal_is_regression(self):
        """EMPTY = 'no answer' and is the WORST state: X → EMPTY is a regression
        (the supplier withdrew their answer), EMPTY → X is an improvement
        (newly answered)."""
        per = {
            "v1.xlsx": {
                "REQ-1": _row("REQ-1", "R1", "NOK"),
                "REQ-2": _row("REQ-2", "R2", "OK"),
                "REQ-3": _row("REQ-3", "R3", "EMPTY"),
            },
            "v2.xlsx": {
                "REQ-1": _row("REQ-1", "R1", "EMPTY"),   # answer removed
                "REQ-2": _row("REQ-2", "R2", "EMPTY"),   # answer removed
                "REQ-3": _row("REQ-3", "R3", "OK"),      # newly answered
            },
        }
        sc, cc, vc, new_in, removed_in, steps = _compute_deltas(
            per, ["v1.xlsx", "v2.xlsx"], {"REQ-1", "REQ-2", "REQ-3"})
        by_id = {c["reqId"]: c for c in sc}
        assert by_id["REQ-1"]["changeType"] == "removed"
        assert by_id["REQ-1"]["improvement"] is False
        assert by_id["REQ-2"]["changeType"] == "removed"
        assert by_id["REQ-3"]["changeType"] == "added"
        assert by_id["REQ-3"]["improvement"] is True
        # completeness: 2 answered -> 1 answered (1 lost)
        assert steps[0]["answeredBefore"] == 2
        assert steps[0]["answeredAfter"] == 1

    def test_new_and_removed(self):
        per = {
            "v1.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK")},
            "v2.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK"),
                        "REQ-2": _row("REQ-2", "R2", "OK")},
        }
        sc, cc, vc, new_in, removed_in, steps = _compute_deltas(
            per, ["v1.xlsx", "v2.xlsx"], {"REQ-1", "REQ-2"})
        assert new_in["1"] == ["REQ-2"]
        assert removed_in["1"] == []

    def test_comment_and_version_changes(self):
        per = {
            "v1.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK", comment="old", version="A")},
            "v2.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK", comment="new", version="B")},
        }
        sc, cc, vc, new_in, removed_in, steps = _compute_deltas(
            per, ["v1.xlsx", "v2.xlsx"], {"REQ-1"})
        assert sc == []
        assert len(cc) == 1 and cc[0]["fromComment"] == "old" and cc[0]["toComment"] == "new"
        assert len(vc) == 1 and vc[0]["fromVersion"] == "A" and vc[0]["toVersion"] == "B"

    def test_three_matrices_two_steps(self):
        per = {
            "v1.xlsx": {"REQ-1": _row("REQ-1", "R1", "NOK")},
            "v2.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK")},
            "v3.xlsx": {"REQ-1": _row("REQ-1", "R1", "NOK")},
        }
        sc, cc, vc, new_in, removed_in, steps = _compute_deltas(
            per, ["v1.xlsx", "v2.xlsx", "v3.xlsx"], {"REQ-1"})
        assert len(steps) == 2
        assert len(sc) == 2
        assert [c["step"] for c in sc] == [1, 2]

    def test_comment_change_detected_when_category_also_changes(self):
        """A supplier can reword a comment while flipping the status — the
        comment change must NOT be swallowed by the status change."""
        per = {
            "v1.xlsx": {"REQ-1": _row("REQ-1", "R1", "NOK", comment="old comment")},
            "v2.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK", comment="new comment")},
        }
        sc, cc, vc, new_in, removed_in, steps = _compute_deltas(
            per, ["v1.xlsx", "v2.xlsx"], {"REQ-1"})
        assert len(sc) == 1          # status changed
        assert len(cc) == 1          # AND the comment changed — both reported
        assert cc[0]["fromComment"] == "old comment"
        assert cc[0]["toComment"] == "new comment"

    def test_formatting_only_comment_diff_not_detected(self):
        """Line-break / whitespace-only differences (ODS keeps \\n, XLSX
        collapses them) must NOT count as a comment change."""
        per = {
            "v1.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK",
                                      comment="line1 line2 line3")},
            "v2.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK",
                                      comment="line1\nline2\nline3")},
        }
        sc, cc, vc, new_in, removed_in, steps = _compute_deltas(
            per, ["v1.xlsx", "v2.xlsx"], {"REQ-1"})
        assert cc == []

    def test_special_char_spacing_diff_not_detected(self):
        """Spacing around measurement symbols (± ≤ ≥ °) and the full-width
        colon (：) is formatting, not content — must not count as a change."""
        per = {
            "v1.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK",
                                      comment="EE： Center±2mm ≤1ms")},
            "v2.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK",
                                      comment="EE：\nCenter ±2mm\n≤ 1ms")},
        }
        sc, cc, vc, new_in, removed_in, steps = _compute_deltas(
            per, ["v1.xlsx", "v2.xlsx"], {"REQ-1"})
        assert cc == []

    def test_genuine_value_change_still_detected(self):
        """A real content change (different measured value) must still be
        detected even when it sits next to a measurement symbol."""
        per = {
            "v1.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK", comment="Center±2mm")},
            "v2.xlsx": {"REQ-1": _row("REQ-1", "R1", "OK", comment="Center±3mm")},
        }
        sc, cc, vc, new_in, removed_in, steps = _compute_deltas(
            per, ["v1.xlsx", "v2.xlsx"], {"REQ-1"})
        assert len(cc) == 1

    def test_new_and_removed_sheet_lists_requirements(self):
        """The 'New & Removed' sheet must list genuinely added and removed
        requirements with their action."""
        from app.qa.conformity_analyzer import MatrixComparison
        c = MatrixComparison()
        c.matrices = [
            {"fileName": "v1.xlsx", "sheetName": "s", "totalRows": 2,
             "stats": {"OK": 1, "NOK": 0, "NA": 0, "EMPTY": 1},
             "inconsistencies": 0, "summary": {"ok": 1, "nok": 0, "na": 0, "empty": 1}},
            {"fileName": "v2.xlsx", "sheetName": "s", "totalRows": 2,
             "stats": {"OK": 2, "NOK": 0, "NA": 0, "EMPTY": 0},
             "inconsistencies": 0, "summary": {"ok": 2, "nok": 0, "na": 0, "empty": 0}},
        ]
        c.new_in = {"1": ["REQ-2"]}
        c.removed_in = {"1": ["REQ-3"]}
        c.steps = [{"step": 1, "matrix1": "v1.xlsx", "matrix2": "v2.xlsx",
                    "statusChanges": 0, "new": 1, "removed": 1,
                    "commentChanges": 0, "versionChanges": 0,
                    "answeredBefore": 1, "answeredAfter": 2}]
        xlsx = generate_delta_excel(c)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(xlsx))
        ws3 = wb["New & Removed"]
        # Row 2 = REQ-2 NEW, Row 3 = REQ-3 REMOVED
        assert ws3.cell(row=2, column=2).value == "REQ-2"
        assert ws3.cell(row=2, column=3).value == "NEW"
        assert ws3.cell(row=3, column=2).value == "REQ-3"
        assert ws3.cell(row=3, column=3).value == "REMOVED"


class TestRealDelta:
    @pytest.fixture(scope="class")
    def comparison(self):
        if not (TIANMA_V1.exists() and TIANMA_V2.exists()):
            pytest.skip("TIANMA matrices not found")
        return compare_matrices(
            [str(TIANMA_V1), str(TIANMA_V2)],
            [TIANMA_V1.name, TIANMA_V2.name],
        )

    def test_delta_fields_present(self, comparison):
        d = comparison_to_dict(comparison)
        for key in ("steps", "newIn", "removedIn", "commentChanges",
                    "versionChanges", "statusChanges", "reportExcel"):
            assert key in d or key == "reportExcel"  # reportExcel added by the route
        assert len(d["steps"]) == 1

    def test_delta_excel_generates(self, comparison):
        xlsx = generate_delta_excel(comparison)
        assert xlsx[:2] == b"PK"
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(xlsx))
        assert "Delta Overview" in wb.sheetnames
        assert "Status Changes" in wb.sheetnames
        assert "New & Removed" in wb.sheetnames

    def test_empty_sheets_show_placeholder_not_bare_header(self):
        """When nothing was added/removed and no comment changed, the sheets
        must say so explicitly instead of looking broken."""
        from app.qa.conformity_analyzer import MatrixComparison
        c = MatrixComparison()
        c.matrices = [
            {"fileName": "v1.xlsx", "sheetName": "s", "totalRows": 1,
             "stats": {"OK": 1, "NOK": 0, "NA": 0, "EMPTY": 0},
             "inconsistencies": 0, "summary": {"ok": 1, "nok": 0, "na": 0, "empty": 0}},
            {"fileName": "v2.xlsx", "sheetName": "s", "totalRows": 1,
             "stats": {"OK": 1, "NOK": 0, "NA": 0, "EMPTY": 0},
             "inconsistencies": 0, "summary": {"ok": 1, "nok": 0, "na": 0, "empty": 0}},
        ]
        c.steps = [{"step": 1, "matrix1": "v1.xlsx", "matrix2": "v2.xlsx",
                    "statusChanges": 0, "new": 0, "removed": 0,
                    "commentChanges": 0, "versionChanges": 0,
                    "answeredBefore": 1, "answeredAfter": 1}]
        xlsx = generate_delta_excel(c)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(xlsx))
        ws3 = wb["New & Removed"]
        assert ws3.cell(row=2, column=2).value and "No requirements" in ws3.cell(row=2, column=2).value
        ws4 = wb["Comment Changes"]
        assert ws4.cell(row=2, column=2).value and "No comment changes" in ws4.cell(row=2, column=2).value