"""
Reference section mapper — extracts section structure from CTS template
and writing guide to enable section→rule citation in findings.

Builds:
- Template section→rules mapping (what each template section requires)
- Guide section→rules mapping (what the writing guide mandates per section)
- Text-only section content analyzer for user documents
"""
import re
from typing import List, Dict, Any, Optional, Tuple
from pathlib import Path
from docx import Document


# ── Known CTS Template section requirements ──────────────────────
# Maps template section names to the rules they enforce
TEMPLATE_SECTION_RULES: Dict[str, List[str]] = {
    "PURPOSE": [
        "Must specify the goal of the specification",
        "Must list and refine requirements impacting component functionality",
        "Must locate the part in its environment via contextual diagram",
        "Must specify performance, operational, constraint, integration, and validation requirements",
    ],
    "SCOPE": [
        "Must define the scope of the specification",
        "Must identify applicable projects and markets (EMEA, North America)",
    ],
    "SYSTEM DEVELOPMENT CONTEXT": [
        "Must describe the requirements applicable to the component",
        "Must reference the generic specification document",
        "Must express stakeholder needs by STELLANTIS system engineering",
    ],
    "GENERAL DESCRIPTION OF THE SYSTEM": [
        "Must present the system purpose",
        "Must include system roles and physical architecture",
    ],
    "SYSTEM ROLES": [
        "Must describe the role of the component in the system",
        "Must specify location and triggering conditions",
        "Must describe operational states and transitions",
    ],
    "PHYSICAL SYSTEM ARCHITECTURE": [
        "Must represent main physical connections between components",
        "Must follow graphic conventions for diagrams",
    ],
    "SYSTEM DIVERSITY": [
        "Must treat complexity by variants",
        "Must specify functional diversity and architecture diversity",
    ],
    "QUOTED DOCUMENTS": [
        "Must list all quoted/referenced documents",
    ],
    "REFERENCE DOCUMENTS": [
        "Must list input specifications for building the TS",
        "Must reference upstream requirements documents",
    ],
    "UPSTREAM REQUIREMENTS": [
        "Must list constraint requirements from other disciplines",
        "Must include regulation and consumerism requirements",
    ],
    "MANDATORY REQUIREMENTS": [
        "Must identify mandatory documents for project management",
        "Must include att_bool@I attribute for mandatory compliance",
    ],
    "APPLICABLE DOCUMENTS": [
        "Must list applicable standards and technical specifications",
    ],
    "STANDARDS": [
        "Must reference applicable standards",
    ],
    "TECHNICAL SPECIFICATIONS": [
        "Must reference technical specifications for connectors",
        "Must reference fault finding and download specifications",
    ],
    "TERMINOLOGY": [
        "Must define glossary, dependability vocabulary, measuring units",
    ],
    "GLOSSARY": [
        "Must define specific vocabulary for the component",
    ],
    "ACRONYMS": [
        "Must list all acronyms used in the document",
    ],
    "REQUIREMENTS": [
        "Must apply requirement engineering template",
        "Functional requirement: 'The system shall <do something> with <performance> in <mode>'",
        "Constructional requirement: 'The system shall <be made of> with <performance> in <configuration>'",
        "Must assign unique requirement IDs following naming convention",
        "Must define ASIL grade where applicable",
        "Complex behavior must be formalized with behavioral models",
    ],
    "FUNCTIONAL REQUIREMENTS": [
        "Must present functional breakdown of requirements",
        "Must not restrict design by the functional breakdown",
        "Must include contextual diagrams and I/O lists",
        "Must define functional states and timing performances",
        "Must specify arming, surveillance, alarm, and disarming states",
    ],
    "PERFORMANCE REQUIREMENTS": [
        "Must specify component performance requirements",
        "Must specify time requirements",
    ],
    "EXTERNAL INTERFACES REQUIREMENTS": [
        "Must include context diagram",
        "Must specify reception and sending frames",
        "Must specify LIN communication rules and physical layers",
        "Must specify heartbeat monitoring",
        "Must define input/output frames and mux tables",
    ],
    "ELECTRICAL INTERFACES": [
        "Must define power supply requirements for all vehicle life situations",
        "Must specify wired connections and connector requirements",
    ],
    "MECHANICAL INTERFACES": [
        "Must specify mechanical interface requirements",
    ],
    "HUMAN-MACHINE INTERFACES": [
        "Must specify HMI requirements or state N/A",
    ],
    "OPERATIONAL REQUIREMENTS": [
        "Must define mission profile and lifetime requirements",
    ],
    "MISSION PROFILE": [
        "Must define mission profile tables",
        "Must specify electronic estimated reliability",
    ],
    "LIFETIME": [
        "Must specify component lifetime requirements",
    ],
    "ERGONOMICS AND HUMAN FACTORS": [
        "Must specify random noise requirements and test methods",
        "Must specify odor requirements",
    ],
    "RAMS REQUIREMENTS": [
        "Must comply with ISO 26262 for safety-critical components",
        "Must define safety requirements and failure mode mitigation",
        "Must specify Technical Safety Requirements (TSR)",
        "Must include SOTIF requirements if applicable",
        "Must specify threat and stress requirements",
        "Must include availability, reliability, durability requirements",
    ],
    "SAFETY REQUIREMENTS": [
        "Must define failure modes contributing to dreaded events",
        "Must associate failure modes with probability requirements",
        "Must list vehicle dreaded events",
    ],
    "MAINTAINABILITY": [
        "Must specify diagnostic requirements",
        "Must define technical interface with diagnostic tools",
        "Must specify self-test procedures and download/remote coding",
        "Must define fault list and strategy",
        "Must specify repair capability and interchangeability",
    ],
    "PRODUCT QUALITY": [
        "Must specify reliability requirements",
        "Must specify quality convergence requirements",
    ],
    "CONSTRAINT REQUIREMENTS": [
        "Must include regulation and consumerism requirements",
        "Must specify weight and physical characteristics",
    ],
    "DESIGN AND MANUFACTURING": [
        "Must specify imposed solutions and design constraints",
        "Must specify materials requirements (% green materials)",
        "Must specify manufacturing requirements",
        "Must specify marking of components",
    ],
    "ENVIRONMENT CONDITIONS": [
        "Must specify temperature resistance requirements",
        "Must specify resistance to acid vapors and vibration",
        "Must specify resistance to electromagnetic interferences",
        "Must specify material behavior to environmental constraints",
        "Must specify impact resistance",
    ],
    # [COMMENTED OUT — Validation requirements rules disabled]
    # "INTEGRATION AND VALIDATION REQUIREMENTS": [
    #     "Must describe evidence requirements for validation",
    #     "Must specify electrical environment, EMC, mechanical, climatic, chemical, and hardware tests",
    #     "Must specify imposed elements of validation plan",
    # ],
    "DEMONSTRATION OF COMPLIANCE WITH REQUIREMENTS": [
        "Must describe supplier proof requirements",
        "Must reference STA19 document for validation evidence",
    ],
    # [COMMENTED OUT — Validation plan rules disabled]
    # "IMPOSED ELEMENTS OF VALIDATION PLAN": [
    #     "Must specify mechanical and climatic test queue",
    #     "Must reference STA19 document",
    #     "Must define test environment and operating modes",
    # ],
    "DOCUMENT": [
        "Must include author, co-author, inspector, and approver information",
        "Must include document identification and project applicability",
        "Must include revision history (index, date, author, modifications)",
        "Must list complexity criteria and variants",
        "Must include applicable documents table with marks and references",
    ],
}

# ── Writing Guide section requirements ───────────────────────────
GUIDE_SECTION_RULES: Dict[str, List[str]] = {
    "PURPOSE": [
        "The Specification is the contractual document between STELLANTIS and supplier",
        "Must contain ALL applicable requirements for the component",
        "Must NOT contain implementation details (how to do)",
        "Must define acceptance criteria for each requirement",
    ],
    "SCOPE": [
        "Applies to all Requirements Documents produced by MECH entity",
        "Covers generic documents and application project documents",
    ],
    "REQUIREMENT WRITING RULES": [
        "Each requirement must have a unique, immutable identifier",
        "Requirements must be verifiable: measurable or testable",
        "Requirements must be unambiguous: one interpretation only",
        "Requirements must be atomic: one requirement per statement",
        "Requirements must include preconditions (WHEN/IF/UNDER)",
        "Requirements must include a trigger (SHALL/MUST)",
        "Requirements must be feasible within project constraints",
        "Do NOT use: 'etc.', 'if possible', 'approximately', 'should', 'may'",
        "Use active voice: 'The system shall...' not 'It is required that...'",
    ],
    "TRACEABILITY": [
        "Each requirement must trace to upstream requirements",
        "Input requirement column must reference parent document and ID",
        "Traceability is mandatory for ISO 26262 compliance",
        "Missing traceability = document rejection",
    ],
    # [COMMENTED OUT — Validation rules in writing guide disabled]
    # "VALIDATION": [
    #     "Each requirement must have a validation method defined",
    #     "Validation method must be specific (test, analysis, inspection, demonstration)",
    #     "Acceptance criteria must be quantified where possible",
    #     "Validation plan must cover all requirements",
    # ],
    "TEMPLATE CLEANLINESS": [
        "ALL placeholders (<<...>>, TBD, XXX) must be removed before submission",
        "Red text/template instructions must be removed",
        "Template examples must be replaced with real content",
        "N/A must be justified, not used as placeholder",
    ],
    "DOCUMENT STRUCTURE": [
        "Must follow the CTS standard plan without deviation",
        "All mandatory sections must be present and completed",
        "Section numbering must follow the CTS standard",
    ],
}


def extract_reference_section_structure(docx_path: str) -> Dict[str, Any]:
    """
    Extract the section structure from a reference document (template or guide).
    Returns a mapping of section_name -> list of content blocks.
    """
    doc = Document(docx_path)
    sections: Dict[str, List[str]] = {}
    current_section = "PREAMBLE"
    section_content: List[str] = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        
        style_name = para.style.name if para.style else ""
        
        # Detect headings via Word styles
        if style_name.startswith("Heading") or style_name.startswith("Titre"):
            try:
                level = int(''.join(filter(str.isdigit, style_name)) or '1')
            except ValueError:
                level = 1
            
            if level <= 2 and section_content:
                # Save previous section
                sections[current_section] = section_content
                section_content = []
            
            if level <= 2:
                current_section = text.upper()
        
        # Also detect ALL-CAPS headings
        if re.match(r'^[A-Z][A-Z\s/()-]{5,}$', text) and len(text) > 6:
            if section_content:
                sections[current_section] = section_content
                section_content = []
            current_section = text.upper()
        
        section_content.append(text)
    
    # Save last section
    if section_content:
        sections[current_section] = section_content
    
    return sections


def get_reference_rules(section_name: str, source: str = "template") -> List[str]:
    """
    Get the CTS rules that apply to a given section.
    
    Args:
        section_name: Name of the section (e.g., "FUNCTIONAL REQUIREMENTS")
        source: "template" or "guide"
    
    Returns:
        List of applicable rules
    """
    section_upper = section_name.upper().strip()
    rules_map = TEMPLATE_SECTION_RULES if source == "template" else GUIDE_SECTION_RULES
    
    # Exact match
    if section_upper in rules_map:
        return rules_map[section_upper]
    
    # Partial match (section name contains or is contained by rule key)
    for key, rules in rules_map.items():
        if key in section_upper or section_upper in key:
            return rules
    
    # Default rules for unknown sections
    return [
        "Must conform to the CTS standard plan",
        "Must contain finalized, project-specific content",
        "Must not contain template placeholders or instructions",
    ]


def analyze_text_section_content(
    section_name: str, 
    section_text: str,
    section_lines: List[str]
) -> Dict[str, Any]:
    """
    Analyze the content of a text-only section for quality issues.
    
    Returns:
        Dict with status, issues found, and content metrics
    """
    result = {
        "section": section_name,
        "status": "ok",
        "issues": [],
        "char_count": len(section_text),
        "line_count": len(section_lines),
        "has_substantive_content": False,
    }
    
    if not section_text.strip():
        result["status"] = "empty"
        result["issues"].append("Section is completely empty — content required per CTS standard plan")
        return result
    
    # Check if section is just "NA" or "N/A"
    na_only = all(
        re.match(r'^(?:N/?A|Not\s+Applicable|NA)\s*$', l.strip(), re.IGNORECASE)
        for l in section_lines if l.strip()
    )
    if na_only:
        result["status"] = "na_only"
        result["issues"].append("Section contains only 'N/A' — verify if this is justified or if content is missing")
        return result
    
    # Check for template placeholders
    placeholder_count = len(re.findall(r'<<[^>]*>>', section_text))
    if placeholder_count > 0:
        result["issues"].append(
            f"Contains {placeholder_count} template placeholder(s) — "
            f"must be replaced with finalized content"
        )
        result["status"] = "has_placeholders"
    
    # Check for TBD/XXX markers
    tbd_count = len(re.findall(r'\bTBD\b', section_text, re.IGNORECASE))
    xxx_count = len(re.findall(r'\bXXX\b', section_text))
    if tbd_count + xxx_count > 0:
        result["issues"].append(
            f"Contains {tbd_count} TBD and {xxx_count} XXX markers — incomplete content"
        )
        if result["status"] == "ok":
            result["status"] = "has_placeholders"
    
    # Check if section has substantive content (not just headings/template)
    # Filter out headings and N/A-only lines, then check TOTAL content
    content_lines = [
        l.strip() for l in section_lines 
        if l.strip() 
        and not re.match(r'^[A-Z][A-Z\s/()-]{5,}$', l.strip())  # Not a heading
        and not re.match(r'^(?:N/?A|Not\s+Applicable)\s*$', l.strip(), re.IGNORECASE)
    ]
    total_content_chars = sum(len(l) for l in content_lines)
    
    # A section has substantive content if:
    # - It has at least 5 non-heading lines, OR
    # - Total content exceeds 100 characters (e.g., a reference table with short cells like "[A1]")
    has_substantive = len(content_lines) >= 5 or total_content_chars >= 100
    result["has_substantive_content"] = has_substantive
    
    if not has_substantive and result["status"] == "ok":
        result["status"] = "minimal_content"
        result["issues"].append(
            "Section has minimal substantive content — may need expansion"
        )
    
    # Check for requirement-like statements (shall/must)
    shall_count = len(re.findall(r'\bshall\b', section_text, re.IGNORECASE))
    must_count = len(re.findall(r'\bmust\b', section_text, re.IGNORECASE))
    if shall_count + must_count > 0:
        result["requirement_statements"] = shall_count + must_count
    
    return result
