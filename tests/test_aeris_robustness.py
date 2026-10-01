"""
AERIS — tests de robustesse, de justesse et de sécurité.

Le fichier `test_aeris_crosscheck.py` prouve que les cas Tianma sont
correctement détectés. Celui-ci prouve les propriétés qui font qu'un
qualiticien peut faire confiance à la liste d'incohérences :

* un fournisseur honnête ne déclenche AUCUNE incohérence (le point le
  plus important : un outil qui crie au loup n'est jamais relu) ;
* une exigence n'emprunte pas les mesures d'une autre ;
* une valeur pile à la limite est jugée correctement ;
* l'unité affichée est celle du document, pas l'unité interne ;
* les formats réellement déposés (.pptx .docx .pdf .txt, .xlsx .xlsm)
  sont lus, y compris quand la bibliothèque principale les refuse ;
* un fichier vide, corrompu ou piégé ne fait jamais tomber le service ;
* le verdict ne dépend pas de la mise en page du dossier.
"""
from __future__ import annotations

import io
import time
import zipfile
from pathlib import Path
from typing import List, Sequence, Tuple

import pytest
from openpyxl import Workbook, load_workbook

from app.qa.aeris_constraints import (
    compare_constraint,
    extract_constraints,
    extract_measurements,
    summarize_verdicts,
)
from app.qa.aeris_crosscheck import report_to_dict, run_crosscheck
from app.qa.aeris_evidence import parse_evidence_bytes
from app.qa.aeris_report import generate_aeris_excel


# ── outils de fixture ─────────────────────────────────────────────

MATRIX_HEADER = [
    "Numéro de l'exigence", "Référence", "", "", "",
    "Libellé", "Conformité FNR", "Commentaires FNR",
]


def write_matrix(path: Path, rows: Sequence[Tuple[str, str, str, str, str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Matrix"
    ws.append(MATRIX_HEADER)
    for req, ref, desc, status, comment in rows:
        ws.append([req, ref, "", "", "", desc, status, comment])
    wb.save(path)
    return path


def run(tmp_path: Path, rows, tdr_text: str, tag: str = "x"):
    matrix = write_matrix(tmp_path / f"matrix_{tag}.xlsx", rows)
    tdr = tmp_path / f"tdr_{tag}.txt"
    tdr.write_text(tdr_text, encoding="utf-8")
    return run_crosscheck(str(matrix), [(tdr.name, tdr.read_bytes())], matrix.name)


def verdict_of(requirement: str, evidence: str) -> str:
    cs = extract_constraints(requirement)
    assert cs, f"aucune contrainte extraite de {requirement!r}"
    verdicts = compare_constraint(cs[0], extract_measurements(evidence))
    return summarize_verdicts(verdicts) if verdicts else "AUCUNE_MESURE"


# ── 1. Aucun faux positif ─────────────────────────────────────────

HONEST_ROWS = [
    ("REQ-1000", "EE", "Current consumption shall be <=100mA", "OK", "Max 82 mA"),
    ("REQ-1001", "OPT", "TFT contrast ratio >=400:1", "OK", "520:1"),
    ("REQ-1002", "ME", "Display temperature <38°C", "OK", "34.2°C"),
    ("REQ-1003", "SW", "CPU load <70%", "OK", "51%"),
    ("REQ-1004", "EE", "Standby current <=20mA", "OK", "12 mA"),
    ("REQ-1005", "ME", "Mechanical gap shall be <=0.3 mm", "OK", "0.18 mm"),
    ("REQ-1006", "OPT", "Luminance shall be >=500 cd/m2", "OK", "640 cd/m2"),
    ("REQ-1007", "SYS", "Startup time <=500ms", "OK", "410 ms"),
    ("REQ-1008", "EE", "Inrush current <=2A", "NOK", "3.1 A — Deviation"),
    ("REQ-1009", "ME", "Weight shall be <=450 g", "NOK", "512 g — Deviation"),
    ("REQ-1010", "SYS", "Diagnostic session not applicable", "NA", "NA"),
]

HONEST_TDR = """Slide 10
EE Power Consumption
Current consumption measured Max 82 mA
Result: OK

Slide 11
Standby
Standby current 12 mA
Result: OK

Slide 12
Inrush
Inrush current 3.1 A NOK Deviation

Slide 20
Optical
TFT contrast ratio 520:1
Result: OK

Slide 21
Luminance
Luminance 640 cd/m2
Result: OK

Slide 30
Thermal
Display temperature 34.2°C
Result: OK

Slide 31
Mechanical
Mechanical gap 0.18 mm
Result: OK
Weight 512 g NOK Deviation

Slide 40
Software
CPU load estimation 51%
Result: OK

Slide 41
Startup
Startup time 410 ms
Result: OK
"""


class TestNoFalsePositives:
    """Un fournisseur honnête ne doit produire aucune incohérence."""

    def test_honest_supplier_raises_nothing(self, tmp_path):
        report = run(tmp_path, HONEST_ROWS, HONEST_TDR, "honest")
        assert report.incoherences == [], (
            "faux positifs : "
            + "; ".join(f"{r['req_id']} {r['motif']}" for r in report.incoherences)
        )

    def test_honest_supplier_verdicts_are_right(self, tmp_path):
        report = run(tmp_path, HONEST_ROWS, HONEST_TDR, "honest2")
        status = {i.req_id: i.final_status for i in report.items}
        for req in ("REQ-1000", "REQ-1001", "REQ-1002", "REQ-1003",
                    "REQ-1004", "REQ-1005", "REQ-1006", "REQ-1007"):
            assert status[req] == "CONFORME", f"{req} -> {status[req]}"
        assert status["REQ-1010"] == "NA"
        # NOK assumé et confirmé par le TDR : une non-conformité déclarée,
        # pas une incohérence.
        assert status["REQ-1008"] in ("NON_CONFORME", "DEVIATION")

    def test_evidence_is_not_borrowed_from_another_requirement(self, tmp_path):
        """≤100 mA ne doit pas absorber les 3.1 A du courant d'appel."""
        report = run(tmp_path, HONEST_ROWS, HONEST_TDR, "isolation")
        item = next(i for i in report.items if i.req_id == "REQ-1000")
        assert "3.1" not in item.supplier_result
        assert "3100" not in item.supplier_result
        assert item.final_status == "CONFORME"

    def test_quantity_label_does_not_leak_across_lines(self):
        """« gap 0.18 mm » au-dessus de « Weight 512 g » reste une longueur."""
        found = extract_measurements(
            "Mechanical gap 0.18 mm\nResult: OK\nWeight 512 g NOK Deviation"
        )
        gap = next(m for m in found if m.unit_family == "length")
        assert gap.quantity != "weight"


# ── 2. Bornes et unités ───────────────────────────────────────────

class TestBoundaries:

    @pytest.mark.parametrize("requirement,evidence,expected", [
        # limite atteinte exactement
        ("Current <=100mA", "Current measured 100 mA", "CONFORME"),
        ("Current <=100mA", "Current measured 100.1 mA", "NON_CONFORME"),
        ("Display temperature <38°C", "Display temperature 38°C", "NON_CONFORME"),
        ("Display temperature <38°C", "Display temperature 37.9°C", "CONFORME"),
        ("Contrast ratio >=400:1", "Contrast 400:1", "CONFORME"),
        ("Contrast ratio >=400:1", "Contrast 399:1", "NON_CONFORME"),
        ("Attenuation >95%", "attenuation 95%", "NON_CONFORME"),
        # unités mixtes
        ("Inrush current <=2A", "Inrush current 3100 mA", "NON_CONFORME"),
        ("Inrush current <=2A", "Inrush current 1500 mA", "CONFORME"),
        ("Current consumption <=100mA", "Current consumption 0.08 A", "CONFORME"),
        ("Current consumption <=100mA", "Current consumption 0.5 A", "NON_CONFORME"),
        ("Startup time <=500ms", "Startup time 2 sec", "NON_CONFORME"),
        ("Gap shall be <=0.3 mm", "Gap measured 180 um", "CONFORME"),
        # pire cas typ/max
        ("Current <=100mA", "Current Typ 85 mA Max 192.7 mA", "NON_CONFORME"),
        ("Current <=100mA", "Current Typ 60 mA Max 90 mA", "CONFORME"),
        # valeurs négatives
        ("Storage temperature shall be >=-40°C", "Storage temperature -45°C", "NON_CONFORME"),
        ("Storage temperature shall be >=-40°C", "Storage temperature -40°C", "CONFORME"),
        # virgule décimale française
        ("Gap shall be <=0,3 mm", "Gap measured 0,25 mm", "CONFORME"),
        ("Gap shall be <=0,3 mm", "Gap measured 0,45 mm", "NON_CONFORME"),
    ])
    def test_verdict(self, requirement, evidence, expected):
        assert verdict_of(requirement, evidence) == expected


class TestUnitDisplay:
    """L'unité lue est celle du document : « ≤2 A », jamais « ≤2000 mA »."""

    @pytest.mark.parametrize("requirement,shown", [
        ("Inrush current <=2A", "2 A"),
        ("Supply voltage >= 1.5 kV", "1.5 kV"),
        ("Boot time <= 2 sec", "2 s"),
        ("Gap <= 300 um", "300 µm"),
        ("Power <= 1.5 W", "1.5 W"),
        ("Clock >= 2 MHz", "2 MHz"),
        ("Current consumption <=100mA", "100 mA"),
        ("Display temperature <38°C", "38 °C"),
        ("Luminance shall be >=500 cd/m2", "500 cd/m²"),
    ])
    def test_display_keeps_the_written_unit(self, requirement, shown):
        assert extract_constraints(requirement)[0].shown() == shown

    def test_gap_is_reported_in_the_requirement_unit(self):
        constraint = extract_constraints("Inrush current <=2A")[0]
        verdict = compare_constraint(
            constraint, extract_measurements("Inrush current 3100 mA")
        )[0]
        assert verdict.measured == "3.1 A"
        assert verdict.gap_unit == "A"
        assert abs(verdict.gap - 1.1) < 1e-9

    @pytest.mark.parametrize("text,family", [
        ("Luminance shall be >=500 cd/m2", "luminance"),
        ("Luminance >= 500 cd/m²", "luminance"),
        ("Luminance >= 500 nits", "luminance"),
        ("Weight shall be <=450 g", "mass"),
        ("Noise <= 35 dB", "acoustic"),
        ("Lifetime >= 10000 h", "time"),
    ])
    def test_families_a_display_supplier_actually_uses(self, text, family):
        assert extract_constraints(text)[0].unit_family == family

    def test_g_force_is_not_a_weight(self):
        """« Vibration up to 50 g » ne doit pas devenir une masse."""
        found = extract_constraints("Vibration up to 50 g")
        assert not any(c.unit_family == "mass" for c in found)


# ── 3. Formats réellement déposés ─────────────────────────────────

def _zip_of(parts: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, body in parts.items():
            zf.writestr(name, body)
    return buf.getvalue()


SLIDE_XML = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<p:cSld><p:spTree><p:sp><p:txBody>'
    '<a:p><a:r><a:t>EE Power Consumption</a:t></a:r></a:p>'
    '<a:p><a:r><a:t>Reduced mode current Typ 119.9 mA Max 192.7 mA NOK</a:t></a:r></a:p>'
    '</p:txBody></p:sp></p:spTree></p:cSld></p:sld>'
)

DOCUMENT_XML = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:body>'
    '<w:p><w:r><w:t>EE Power Consumption</w:t></w:r></w:p>'
    '<w:p><w:r><w:t>Reduced mode current Typ 119.9 mA Max 192.7 mA NOK</w:t></w:r></w:p>'
    '</w:body></w:document>'
)


class TestEvidenceFormats:

    def test_pptx_written_by_python_pptx(self, tmp_path):
        pptx = pytest.importorskip("pptx")
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "EE Power Consumption"
        box = slide.shapes.add_textbox(
            pptx.util.Inches(1), pptx.util.Inches(2),
            pptx.util.Inches(6), pptx.util.Inches(1),
        )
        box.text_frame.text = "Reduced mode current Typ 119.9 mA Max 192.7 mA NOK"
        slide.notes_slide.notes_text_frame.text = "Startup time 450 ms"
        path = tmp_path / "real.pptx"
        prs.save(path)

        doc = parse_evidence_bytes("real.pptx", path.read_bytes())
        joined = " ".join(c.text for c in doc.chunks)
        assert doc.chunks
        assert "119.9" in joined
        assert "450 ms" in joined          # les notes comptent comme preuve

    def test_pptx_rejected_by_python_pptx_falls_back_to_xml(self):
        """
        Un PPTX exporté par un autre outil fait échouer python-pptx. Sans
        repli, tout le TDR serait lu vide et chaque exigence passerait en
        « aucune preuve ».
        """
        content = _zip_of({"ppt/slides/slide1.xml": SLIDE_XML})
        doc = parse_evidence_bytes("exported.pptx", content)
        assert doc.chunks, "le lecteur de secours n'a pas pris le relais"
        assert "119.9" in doc.chunks[0].text
        assert doc.chunks[0].location == "Slide 1"

    def test_docx_falls_back_to_xml(self):
        content = _zip_of({"word/document.xml": DOCUMENT_XML})
        doc = parse_evidence_bytes("exported.docx", content)
        assert doc.chunks
        assert "119.9" in " ".join(c.text for c in doc.chunks)

    def test_txt_keeps_slide_locations(self):
        doc = parse_evidence_bytes("tdr.txt", HONEST_TDR.encode("utf-8"))
        assert {"Slide 10", "Slide 20", "Slide 41"} <= {c.location for c in doc.chunks}

    @pytest.mark.parametrize("name,content", [
        ("vide.txt", b""),
        ("vide.pptx", b""),
        ("corrompu.pptx", b"ceci n'est pas un zip"),
        ("corrompu.pdf", b"%PDF-1.4 tronque"),
        ("corrompu.docx", b"PK\x03\x04 mais pas un docx"),
        ("renomme.pptx", b"%PDF-1.4\n%%EOF\n"),
        ("binaire.txt", bytes(range(256)) * 10),
        ("utf16.txt", "Slide 1\nCurrent 42 mA".encode("utf-16")),
        ("latin1.txt", "Slide 1\nTempérature 38°C".encode("latin-1")),
    ])
    def test_damaged_file_never_raises(self, name, content):
        doc = parse_evidence_bytes(name, content)
        assert isinstance(doc.chunks, list)


class TestMatrixFormats:

    @pytest.mark.parametrize("name", ["matrix.xlsx", "matrix.xlsm"])
    def test_excel_variants(self, tmp_path, name):
        from app.qa.conformity_analyzer import extract_conformity_data
        path = write_matrix(
            tmp_path / name,
            [("REQ-0308287", "EE", "Current consumption <=100mA", "NOK", "Typ 119.9 mA")],
        )
        analysis = extract_conformity_data(str(path), name)
        assert any(i.req_id == "REQ-0308287" for i in analysis.items)

    def test_ods_matrix(self, tmp_path):
        pytest.importorskip("odf")
        from odf.opendocument import OpenDocumentSpreadsheet
        from odf.table import Table, TableCell, TableRow
        from odf.text import P
        from app.qa.conformity_analyzer import extract_conformity_data

        doc = OpenDocumentSpreadsheet()
        table = Table(name="Matrix")
        rows = [
            MATRIX_HEADER,
            ["REQ-0308287", "EE", "", "", "",
             "Current consumption <=100mA", "NOK", "Typ 119.9 mA"],
        ]
        for values in rows:
            tr = TableRow()
            for value in values:
                cell = TableCell(valuetype="string")
                cell.addElement(P(text=str(value)))
                tr.addElement(cell)
            table.addElement(tr)
        doc.spreadsheet.addElement(table)
        path = tmp_path / "matrix.ods"
        doc.save(str(path))

        analysis = extract_conformity_data(str(path), "matrix.ods")
        assert any(i.req_id == "REQ-0308287" for i in analysis.items)


# ── 4. Entrées dégradées ──────────────────────────────────────────

class TestDegradedInput:

    def test_matrix_with_headers_but_no_requirement(self, tmp_path):
        """Une feuille sans ligne de données ne doit pas interrompre l'analyse."""
        matrix = write_matrix(tmp_path / "vide.xlsx", [])
        tdr = tmp_path / "tdr.txt"
        tdr.write_text("Slide 1\nCurrent 42 mA", encoding="utf-8")
        report = run_crosscheck(str(matrix), [(tdr.name, tdr.read_bytes())], matrix.name)
        assert report.items == []
        assert report.incoherences == []
        load_workbook(io.BytesIO(generate_aeris_excel(report_to_dict(report))))

    def test_no_evidence_at_all(self, tmp_path):
        matrix = write_matrix(
            tmp_path / "m.xlsx",
            [("REQ-1", "EE", "Current <=100mA", "OK", "80 mA")],
        )
        report = run_crosscheck(str(matrix), [], matrix.name)
        assert report.items[0].final_status == "MANQUANT"

    def test_rows_without_id_or_with_huge_text(self, tmp_path):
        rows = [
            ("", "EE", "Current <=100mA", "OK", "80 mA"),
            ("REQ-X", "", "", "", ""),
            ("REQ-Y", "EE", "A" * 9000, "OK", "B" * 9000),
            ("REQ-Z", "EE", "Current <=100mA", "statut-inconnu", "80 mA"),
        ]
        report = run(tmp_path, rows, "Slide 1\nCurrent 42 mA", "degraded")
        load_workbook(io.BytesIO(generate_aeris_excel(report_to_dict(report))))


class TestExcelSafety:
    """Le texte vient d'un fournisseur externe : il ne doit jamais s'exécuter."""

    def test_formulas_are_neutralised(self, tmp_path):
        rows = [
            ("REQ-E1", "EE", "Current <=100mA", "OK", "=cmd|'/c calc'!A1"),
            ("REQ-E2", "EE", "@SUM(1+1)*cmd", "OK", "+1+1"),
            ("REQ-E3", "EE", "Storage temperature >=-40°C", "OK", "-40 °C"),
        ]
        report = run(tmp_path, rows, "Slide 1\nStorage temperature -40 °C", "evil")
        workbook = load_workbook(
            io.BytesIO(generate_aeris_excel(report_to_dict(report)))
        )
        unprotected: List[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str):
                        continue
                    if value[:1] not in ("=", "+", "-", "@", "\t", "\r"):
                        continue
                    if value[:1] in "+-" and value[1:2].isdigit():
                        continue      # « -40 °C » : un nombre, pas une formule
                    if not cell._style.quotePrefix:
                        unprotected.append(f"{sheet.title}!{cell.coordinate}={value!r}")
        assert not unprotected, f"formules exécutables : {unprotected}"

    def test_negative_values_stay_readable(self, tmp_path):
        rows = [("REQ-N", "EE", "Storage temperature >=-40°C", "OK", "-40 °C")]
        report = run(tmp_path, rows, "Slide 1\nStorage temperature -40 °C", "neg")
        item = report.items[0]
        assert item.target.startswith("≥-40") or "-40" in item.target


# ── 5. Le verdict ne dépend pas de la mise en page ────────────────

BASE_ROWS = [
    ("REQ-2001", "EE", "Current consumption shall be <=100mA", "OK", "Max 192.7 mA"),
    ("REQ-2002", "OPT", "TFT contrast ratio >=400:1", "OK", "500@25C / 380@85C"),
    ("REQ-2003", "ME", "Display temperature <38°C", "OK", "41.6°C"),
    ("REQ-2004", "SW", "CPU load <70%", "OK", "50%"),
    ("REQ-2005", "EE", "Standby current <=20mA", "NOK", "35 mA Deviation"),
]

BASE_TDR = """Slide 10
EE Power Consumption
Current consumption Typ 119.9 mA Max 192.7 mA NOK Deviation

Slide 20
Optical
TFT contrast 500 @25°C  410 @70°C  380 @85°C

Slide 30
Thermal
Display temperature 41.6°C Deviation

Slide 40
Software
CPU load estimation 50% OK

Slide 50
Standby
Standby current 35 mA NOK
"""

NOISE = """
Slide 90
Project organisation
The team met on 2026-04-27 to review the planning.

Slide 91
Packaging
Boxes are 300 mm wide and contain 24 units.
Pallet weight 512 kg. Shipping lead time 45 days.
"""


def fingerprint(report):
    return (
        {i.req_id: i.final_status for i in report.items},
        {r["req_id"]: r["motif"] for r in report.incoherences},
    )


class TestLayoutInvariance:

    @pytest.fixture
    def reference(self, tmp_path):
        return fingerprint(run(tmp_path, BASE_ROWS, BASE_TDR, "ref"))

    def test_requirement_order_does_not_matter(self, tmp_path, reference):
        got = fingerprint(run(tmp_path, BASE_ROWS[::-1], BASE_TDR, "rev"))
        assert got == reference

    def test_slide_order_does_not_matter(self, tmp_path, reference):
        blocks = BASE_TDR.strip().split("\n\n")
        shuffled = "\n\n".join(reversed(blocks))
        assert fingerprint(run(tmp_path, BASE_ROWS, shuffled, "sl")) == reference

    def test_unrelated_slides_do_not_change_anything(self, tmp_path, reference):
        assert fingerprint(run(tmp_path, BASE_ROWS, BASE_TDR + NOISE, "noise")) == reference

    def test_duplicated_tdr_does_not_change_anything(self, tmp_path, reference):
        doubled = BASE_TDR + "\n" + BASE_TDR
        assert fingerprint(run(tmp_path, BASE_ROWS, doubled, "dup")) == reference

    def test_case_and_spacing_do_not_matter(self, tmp_path, reference):
        noisy = BASE_TDR.replace("mA", " MA ").replace("Slide", "SLIDE")
        assert fingerprint(run(tmp_path, BASE_ROWS, noisy, "case")) == reference

    def test_same_value_written_in_another_unit(self, tmp_path, reference):
        """192.7 mA écrit 0.1927 A doit donner le même verdict."""
        rewritten = BASE_TDR.replace(
            "Typ 119.9 mA Max 192.7 mA", "Typ 0.1199 A Max 0.1927 A"
        )
        statuses, _ = fingerprint(run(tmp_path, BASE_ROWS, rewritten, "amp"))
        assert statuses["REQ-2001"] == reference[0]["REQ-2001"]

    def test_runs_are_deterministic(self, tmp_path):
        runs = [fingerprint(run(tmp_path, BASE_ROWS, BASE_TDR, f"d{n}")) for n in range(3)]
        assert runs[0] == runs[1] == runs[2]


class TestMonotonicity:

    def test_a_worse_measurement_cannot_improve_the_verdict(self, tmp_path):
        worse = BASE_TDR.replace("Display temperature 41.6°C", "Display temperature 60°C")
        statuses, _ = fingerprint(run(tmp_path, BASE_ROWS, worse, "worse"))
        assert statuses["REQ-2003"] == "NON_CONFORME"

    def test_a_better_measurement_still_flags_the_stale_comment(self, tmp_path):
        """
        Le TDR passe à 30 °C mais la matrice annonce toujours 41.6 °C :
        l'exigence devient conforme et l'écart entre les deux documents
        du fournisseur reste signalé.
        """
        better = BASE_TDR.replace("Display temperature 41.6°C", "Display temperature 30°C")
        statuses, motifs = fingerprint(run(tmp_path, BASE_ROWS, better, "better"))
        assert statuses["REQ-2003"] == "CONFORME"
        assert "chiffres" in motifs["REQ-2003"]


class TestNegativeControl:

    def test_a_tdr_for_another_product_proves_nothing(self, tmp_path):
        other = (
            "Slide 1\nBrake pedal assembly\nPedal travel 120 mm, return force 45 N\n"
            "\nSlide 2\nHydraulic pressure 180 bar at 20 degrees\n"
        )
        report = run(tmp_path, BASE_ROWS, other, "other")
        statuses = {i.final_status for i in report.items}
        assert "CONFORME" not in statuses, "conformité inventée à partir d'un dossier hors sujet"


# ── 6. Taille réelle ──────────────────────────────────────────────

class TestScale:

    @pytest.mark.slow
    def test_450_requirements_stay_quiet_and_fast(self, tmp_path):
        """La vraie matrice Stellantis dépasse 400 exigences."""
        rows, slides = [], []
        for n in range(450):
            req = f"REQ-{3000 + n}"
            kind = n % 5
            if kind == 0:
                limit, measured = 80 + n % 40, 60 + n % 15
                rows.append((req, "EE",
                             f"Current consumption shall be <={limit}mA", "OK",
                             f"{measured} mA"))
                slides.append(f"Slide {n}\nPower {req}\n"
                              f"Current consumption {measured} mA OK")
            elif kind == 1:
                rows.append((req, "OPT", f"Contrast ratio >={300 + n % 100}:1", "OK", "500:1"))
                slides.append(f"Slide {n}\nOptical {req}\nContrast 500:1")
            elif kind == 2:
                rows.append((req, "ME", f"Temperature <{40 + n % 20}°C", "OK", "35°C"))
                slides.append(f"Slide {n}\nThermal {req}\nTemperature 35°C")
            elif kind == 3:
                rows.append((req, "SW", "CPU load <70%", "NOK", "88% Deviation"))
                slides.append(f"Slide {n}\nSoftware {req}\nCPU load 88% NOK")
            else:
                rows.append((req, "SYS", "Not applicable on this variant", "NA", "NA"))

        started = time.time()
        report = run(tmp_path, rows, "\n\n".join(slides), "big")
        elapsed = time.time() - started

        assert len(report.items) == 450
        assert report.incoherences == [], (
            "faux positifs à l'échelle : "
            + "; ".join(f"{r['req_id']} {r['motif']}" for r in report.incoherences[:5])
        )
        assert elapsed < 120, f"trop lent : {elapsed:.0f}s pour 450 exigences"
