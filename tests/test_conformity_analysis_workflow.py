"""Local end-to-end matrix, reports and dataset tests using synthetic data."""
import base64
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.conformity_server import app
from app.qa import conformity_analyzer as ca
from app.qa.conformity_report import generate_powerbi_dataset


def matrix_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthetic matrix"
    ws.append(["Requirement", "Description", "Conformity FNR", "Comments FNR"])
    for index, (status, comment) in enumerate([
        ("OK", "Thermal simulation is needed"),
        ("OK", "No defect or failure found; no deviation required."),
        ("OK", ""),
        ("NOK", "Test failed"),
        ("DEVIATION", "Approved waiver W-01"),
        ("NA", "Not applicable to this variant"),
        ("", ""),
    ], 1):
        ws.append([f"REQ-{index:05}", "The device shall operate at 85 C.", status, comment])
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(ca, "_analyze_ok_deep_llm", lambda items: ([], set()))
    monkeypatch.chdir(tmp_path)
    return TestClient(app)


def analyze(client):
    response = client.post("/api/conformity-excel", files={
        "file": ("synthetic.xlsx", matrix_bytes(),
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    })
    assert response.status_code == 200, response.text
    return response.json()


def test_single_upload_preserves_statuses_and_grounded_export(client):
    result = analyze(client)
    analysis = result["analysis"]
    assert analysis["summary"]["total"] == 7
    assert analysis["stats"] == {"OK": 3, "NOK": 1, "DEVIATION": 1, "NA": 1, "EMPTY": 1}
    assert sum(analysis["stats"].values()) == len(analysis["items"])
    assert len(analysis["okDeepFindings"]) == 1
    finding = analysis["okDeepFindings"][0]
    assert finding["reqId"] == "REQ-00001"
    assert finding["findingType"] == "pending_verification"
    assert finding["evidenceExcerpt"] == "Thermal simulation is needed"
    assert finding["nextAction"]
    assert analysis["reviewCoverage"]["aiReviewedItems"] == 0
    assert "Review coverage:" in analysis["reportText"]
    wb = load_workbook(BytesIO(base64.b64decode(result["reportExcel"])), data_only=True)
    assert "Review Coverage" in wb.sheetnames
    exported = wb["Deep-Dive OK Analysis"].cell(2, 6).value
    assert finding["nextAction"] in exported
    assert "85 C" in exported
    assert "Source row 2" in exported


def test_batch_partial_failure_and_coverage_are_exported(client):
    response = client.post("/api/conformity-batch", files=[
        ("files", ("synthetic.xlsx", matrix_bytes(), "application/octet-stream")),
        ("files", ("unsupported.csv", b"invalid", "text/csv")),
    ])
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["filesAnalyzed"] == 1
    assert result["filesFailed"] == 1
    assert result["files"][0]["reviewCoverage"]["patternReviewedItems"] == 2
    wb = load_workbook(BytesIO(base64.b64decode(result["reportExcel"])), data_only=True)
    assert "Review Coverage" in wb.sheetnames
    assert wb["Review Coverage"].cell(2, 4).value == 0
    assert wb["Overview"].cell(3, 11).value == "Deviations"
    assert wb["Overview"].cell(4, 11).value == 1


def test_pdf_has_coverage_and_action(client):
    response = client.post("/api/conformity-report", files={
        "file": ("synthetic.xlsx", matrix_bytes(), "application/octet-stream"),
    })
    assert response.status_code == 200, response.text
    result = response.json()
    from PyPDF2 import PdfReader
    pdf_bytes = base64.b64decode(result["reportPdf"])
    assert pdf_bytes.startswith(b"%PDF")
    text = " ".join(" ".join(page.extract_text() for page in PdfReader(BytesIO(pdf_bytes)).pages).split())
    assert "Review coverage:" in text
    assert "Next action:" in text
    assert "not proof of non-compliance" in text
    assert "DECLARED DEVIATIONS" in text
    assert "UNANSWERED REQUIREMENTS" in text
    assert "All OK items are consistent" not in text


def test_powerbi_counts_and_schema_match_exported_rows(client):
    dataset = generate_powerbi_dataset(analyze(client)["analysis"])
    assert sum(row["Count"] for row in dataset["data"]["StatusSummary"]) == 7
    assert any(row["Status"] == "DEVIATION" for row in dataset["data"]["StatusSummary"])
    for table in dataset["dataset"]["tables"]:
        expected = {column["name"] for column in table["columns"]}
        for row in dataset["data"][table["name"]]:
            assert set(row) == expected
    assert dataset["data"]["Inconsistencies"][0]["Explanation"]
    assert dataset["reviewCoverage"]["limitations"]


@pytest.mark.parametrize(("name", "content", "status"), [
    ("empty.xlsx", b"", 400), ("broken.xlsx", b"not a workbook", 422),
    ("invalid.txt", b"text", 400),
])
def test_invalid_upload_is_explicitly_rejected(client, name, content, status):
    response = client.post("/api/conformity-excel", files={
        "file": (name, content, "application/octet-stream"),
    })
    assert response.status_code == status
    assert response.json()["detail"]
