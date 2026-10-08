"""DIA field-role and placeholder regressions, without external model calls."""
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.qa import conformity_analyzer as ca
from app.qa.conformity_report import generate_conformity_excel, generate_batch_conformity_excel


REAL_DIA = Path(
    r"C:\Users\TA29225\Desktop\conformity matrix\TIANMA dossier consultation"
    r"\02016_13_02411_Template_DIA_ISO26262_V6.2-0_Supplier.xlsx"
)


def make_dia(tmp_path, offset=0):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "DIA"
    ws.append([""] * offset + [
        "", "", "", "Development Interface Agreement", "", "", "", "", "", "", "", "", "", "",
    ])
    ws.append([""] * offset + [
        "", "", "", "", "<Select>", "", "", "", "",
        "Supplier Concurrence/Methods/Comments", "", "", "Stellantis Comments",
        "Global agreement status",
    ])
    ws.append([""] * offset + [
        "", "ID", "Reference : ISO 26262:2018", "Stellantis Policy & Procedure",
        "Work Product Delivery Type", "Deliverable", "Stellantis", "<Supplier Name>",
        'Primary "Responsible (Role) Name"',
        "Supplier Concurrence to Stellantis Policy & Procedures",
        'If "No" or "Partial", please provide alternative method and justification',
        "Supplier Assumptions and Comments", "Stellantis Assumptions and comments", "",
    ])
    cases = [
        ("Yes", "", "Provide report after verification", "Customer comment", "<select>"),
        ("Yes", "", "Internal audit is planned", "Not a supplier response", "<Select>"),
        ("No", "Alternative audit method", "Cannot provide summary", "OK", "<select>"),
        ("Not Applicable", "Not used in this variant", "", "Customer opinion", "<select>"),
        ("<select>", "", "<select>", "", "<select>"),
    ]
    for index, (status, alternative, comment, customer, agreement) in enumerate(cases, 1):
        ws.append([""] * offset + [
            "", f"DIA-00-{index:02}", "ISO 26262",
            "The supplier shall provide verification reports.", "2 - Summary",
            "Verification report", "A", "R", "FSM", status, alternative,
            comment, customer, agreement,
        ])
    path = tmp_path / "dia.xlsx"
    workbook.save(path)
    return path


@pytest.mark.parametrize("offset", [0, 1, 3])
def test_dia_uses_supplier_concurrence_and_both_supplier_comment_fields(tmp_path, offset):
    analysis = ca.extract_conformity_data(str(make_dia(tmp_path, offset)))
    assert analysis.header_row == 2
    assert analysis.column_mapping["conformity"] == [offset + 9]
    assert analysis.column_mapping["comment"] == [offset + 10, offset + 11]
    assert analysis.column_mapping["description"] == [offset + 3]
    assert len(analysis.items) == 5
    assert analysis.items[0].comment == "Provide report after verification"
    assert analysis.items[2].comment == "Alternative audit method | Cannot provide summary"
    assert analysis.items[2].conformity_category == "NOK"
    assert analysis.items[4].conformity_category == "EMPTY"
    assert analysis.items[4].comment == ""
    assert all("<select>" not in item.comment.lower() for item in analysis.items)


def test_placeholder_in_comment_is_not_evidence(tmp_path, monkeypatch):
    analysis = ca.extract_conformity_data(str(make_dia(tmp_path)))
    analysis.items[0].comment = "<select>"
    captured = []
    monkeypatch.setattr(ca, "_analyze_ok_deep_llm",
                        lambda items: (captured.extend(items) or [], set()))
    ca.analyze_ok_deep(analysis)
    assert all(it.comment.lower() != "<select>" for it in captured)
    assert analysis.review_coverage["withoutComment"] == 1


def test_placeholder_does_not_infer_a_comment_column():
    rows = [["ID", "Status", "Unnamed"], *[
        [f"DIA-00-{index:02}", "Yes", "<select>"] for index in range(1, 8)
    ]]
    statuses, comments = ca._detect_conformity_columns_by_content(rows, 1)
    assert statuses == [1]
    assert comments == []


@pytest.mark.skipif(not REAL_DIA.is_file(), reason="Local supplier DIA workbook is unavailable")
def test_real_dia_comments_match_source_and_both_excel_exports(monkeypatch):
    monkeypatch.setattr(ca, "_analyze_ok_deep_llm", lambda items: ([], set()))
    analysis = ca.analyze_conformity_matrix(str(REAL_DIA), REAL_DIA.name)
    source = load_workbook(REAL_DIA, data_only=True)["DIA V6.2"]
    assert analysis.sheet_name == "DIA V6.2"
    assert analysis.header_row == 2
    assert len(analysis.items) == 56
    assert analysis.column_mapping["comment"] == [10, 11]
    for item in analysis.items:
        row = item.row_index + 1
        expected = " | ".join(
            str(source.cell(row, col).value).strip()
            for col in (11, 12) if source.cell(row, col).value
        )
        assert item.comment == expected
        assert "<select>" not in item.comment.lower()
        description = source.cell(row, 4).value
        if description is None:
            merged = next((r for r in source.merged_cells.ranges
                           if r.min_col <= 4 <= r.max_col and r.min_row <= row <= r.max_row), None)
            if merged:
                description = source.cell(merged.min_row, merged.min_col).value
        assert item.description == (str(description).strip() if description else "")
    assert "online audit" in next(it.comment for it in analysis.items if it.req_id == "DIA-00-02")
    result = ca.analysis_to_dict(analysis)
    for report in (generate_conformity_excel(result), generate_batch_conformity_excel([result])):
        wb = load_workbook(BytesIO(report), data_only=True)
        ws = next(sheet for sheet in wb if sheet.title == "All Items"
                  or sheet.title.endswith("Items"))
        exported_comments = [ws.cell(row, 6).value or "" for row in range(2, ws.max_row + 1)]
        assert exported_comments == [item.comment for item in analysis.items]
