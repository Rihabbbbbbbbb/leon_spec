"""Extract traceable text and tables from PowerPoint presentations.

This module intentionally makes no compliance decision. It retains slide and
shape/table provenance so a reviewer can inspect every candidate passage.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, List, Sequence

from app.qa.conformity_coverage import extract_id_tokens


@dataclass
class SlideEvidence:
    file_name: str
    slide_number: int
    shape_name: str
    shape_type: str
    text: str
    evidence_ids: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


def extract_pptx_evidence(filepath: str | Path) -> List[SlideEvidence]:
    """Extract visible text and table cell contents from a .pptx.

    Notes, embedded images, charts' underlying workbook data, and legacy .ppt
    files are intentionally not interpreted. Their absence is not evidence of
    non-compliance.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    path = Path(filepath)
    if path.suffix.lower() != ".pptx":
        raise ValueError("Only .pptx presentations are supported; convert legacy .ppt files first.")
    if not path.is_file():
        raise FileNotFoundError(str(path))

    presentation = Presentation(str(path))
    evidence: List[SlideEvidence] = []

    def add_text(slide_number: int, shape_name: str, shape_type: str, text: str) -> None:
        normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
        if not normalized:
            return
        evidence.append(SlideEvidence(
            file_name=path.name,
            slide_number=slide_number,
            shape_name=shape_name or "(unnamed)",
            shape_type=shape_type,
            text=normalized,
            evidence_ids=extract_id_tokens(normalized),
        ))

    def walk_shapes(shapes: Sequence[Any], slide_number: int) -> None:
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk_shapes(shape.shapes, slide_number)
                continue
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    cells = [" ".join(cell.text.split()) for cell in row.cells]
                    if any(cells):
                        rows.append(" | ".join(cells))
                add_text(slide_number, shape.name, "table", "\n".join(rows))
            if getattr(shape, "has_text_frame", False):
                add_text(slide_number, shape.name, "text", shape.text_frame.text)

    for slide_number, slide in enumerate(presentation.slides, start=1):
        walk_shapes(slide.shapes, slide_number)
    return evidence

