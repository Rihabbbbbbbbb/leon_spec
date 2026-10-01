from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.qa.conformity_evidence_report import generate_evidence_excel


def test_evidence_report_shows_empty_requirements_even_without_candidates():
    analysis = {
        "matrixFile": "supplier.xlsm",
        "tdrFile": "report.pdf",
        "summary": {
            "requirements": 1,
            "requirementsWithEvidenceCandidates": 0,
            "requirementsWithoutEvidenceCandidates": 1,
            "pagesWithExtractedText": 4,
            "evidenceBlocks": 0,
        },
        "decisionSummary": {"supplierDeclared": {"OK": 0, "NOK": 0, "NA": 0, "EMPTY": 1}},
        "requirements": [{
            "matrixRowNumber": 42,
            "reqId": "REQ-42",
            "reference": "REF-42",
            "description": "Requirement without a text candidate",
            "matrixStatus": "EMPTY",
            "matrixStatusRaw": "",
            "supplierComment": "",
            "reviewStatus": "no_text_candidate_found",
            "evidence": [],
        }],
        "limitations": ["Scanned content was not assessed."],
    }

    workbook = load_workbook(BytesIO(generate_evidence_excel(analysis)), data_only=True)

    assert workbook["Review Summary"]["A7"].value == "Without text candidates"
    assert workbook["Review Summary"]["B8"].value == 4
    assert workbook["Requirements"]["A2"].value == 42
    assert workbook["Requirements"]["E2"].value == "EMPTY"
    assert workbook["Requirements"]["H2"].value == "no_text_candidate_found"
    assert workbook["Requirements"]["I2"].value == 0


def test_report_marks_candidate_score_as_not_confidence_and_keeps_citation():
    analysis = {
        "matrixFile": "supplier.xlsm",
        "presentationFile": "slides.pptx",
        "summary": {"requirements": 1, "requirementsWithEvidenceCandidates": 1,
                    "requirementsWithoutEvidenceCandidates": 0, "slidesWithExtractedText": 2},
        "requirements": [{
            "matrixRowNumber": 3, "reqId": "REQ-3", "description": "Input shall be 12 V",
            "matrixStatus": "OK", "matrixStatusRaw": "OK", "supplierComment": "Declared OK",
            "reviewStatus": "candidate_evidence_found",
            "evidence": [{"matchType": "text_similarity", "matchScore": 0.25,
                          "file_name": "slides.pptx", "slide_number": 7,
                          "shape_name": "Test results", "text": "Measured input at 12 V"}],
        }],
    }

    workbook = load_workbook(BytesIO(generate_evidence_excel(analysis)), data_only=True)
    sheet = workbook["Evidence Candidates"]
    assert "not confidence" in sheet["E1"].value.lower()
    assert sheet["F2"].value == "slides.pptx"
    assert sheet["G2"].value == 7
    assert sheet["H2"].value == "Test results"
