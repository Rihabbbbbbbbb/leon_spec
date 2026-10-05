"""
Evidence document parser for AERIS matrix↔TDR cross-check.

Reads supplier technical proof (TDR / PPT / PDF / DOCX / TXT) into
page-or-slide-aware chunks. Failed parsing is exposed in document metadata;
an evidence set with no usable text cannot produce a successful synthesis.
"""
from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import List
from xml.etree import ElementTree as ET


_PPT_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

_EVIDENCE_EXT = {".pptx", ".pdf", ".docx", ".txt"}


@dataclass
class EvidenceChunk:
    """One retrievable passage of supplier evidence."""
    file_name: str
    location: str          # "Slide 12" / "Page 90" / "Section …"
    text: str
    chunk_id: int = 0
    kind: str = "slide"    # slide | page | section | paragraph
    measurement_text: str | None = None


@dataclass
class EvidenceDocument:
    file_name: str
    kind: str              # pptx | pdf | docx | txt
    chunks: List[EvidenceChunk] = field(default_factory=list)
    page_count: int = 0
    parse_error: str = ""
    warnings: List[str] = field(default_factory=list)


def is_evidence_file(name: str) -> bool:
    return Path(name).suffix.lower() in _EVIDENCE_EXT


def parse_evidence_bytes(file_name: str, content: bytes) -> EvidenceDocument:
    """Parse an uploaded evidence file from memory."""
    suffix = Path(file_name).suffix.lower()
    doc = EvidenceDocument(file_name=file_name, kind=suffix.lstrip(".") or "txt")
    try:
        if suffix == ".ppt":
            raise ValueError("Legacy .ppt is not supported; export it as .pptx or PDF")
        if suffix not in _EVIDENCE_EXT:
            raise ValueError(f"Unsupported evidence format: {suffix or 'none'}")
        if not content:
            raise ValueError("The evidence file is empty")
        if suffix == ".pptx":
            chunks = _extract_pptx(content, file_name, doc.warnings)
            doc.kind = "pptx"
        elif suffix == ".pdf":
            chunks = _extract_pdf(content, file_name, doc.warnings)
            doc.kind = "pdf"
        elif suffix == ".docx":
            chunks = _extract_docx(content, file_name)
            doc.kind = "docx"
        else:
            text = content.decode("utf-8-sig")
            chunks = _chunk_plain(text, file_name)
            doc.kind = "txt"
        doc.chunks = chunks
        if not chunks:
            doc.warnings.append("No usable text was extracted from this evidence document")
        doc.page_count = _count_locations(chunks)
        doc.warnings = list(dict.fromkeys(doc.warnings))
    except Exception as exc:
        doc.parse_error = str(exc)
        doc.chunks = []
    return doc


def parse_evidence_path(path: Path) -> EvidenceDocument:
    return parse_evidence_bytes(path.name, path.read_bytes())


def _count_locations(chunks: List[EvidenceChunk]) -> int:
    return len({c.location for c in chunks})


_DECORATION = re.compile(
    r"(?:©|copyright|all rights reserved|tous droits r[ée]serv[ée]s)", re.I
)
_MAX_IMAGES_PER_PAGE = 20
_MAX_IMAGE_BYTES = 8_000_000


def _table_row_blocks(
    rows: List[List[str | None]],
) -> List[tuple[str, str | None]]:
    """Keep supplier cells separate from the requirement copied in the table."""
    if not rows:
        return []
    header_index = next((
        i for i, row in enumerate(rows[:5])
        if any(re.search(r"req(?:uirement)?(?:\s+(?:id|reference))|description|measurement", cell or "", re.I)
               for cell in row)
        and sum(bool(cell) for cell in row) >= 2
    ), 0)
    headers = [re.sub(r"\s+", " ", cell or "").strip() for cell in rows[header_index]]
    if len(rows) == 1 and len(re.findall(r"\b(?:REQ|REF)-[\w.-]+", " ".join(headers))) > 1:
        return []
    supplier_cols = [
        i for i, header in enumerate(headers)
        if re.search(r"(?:(?:tm|supplier)\s+(?:proposal|offer|result)|measured|measurement result|^result$)", header, re.I)
    ]
    context_cols = [
        i for i, header in enumerate(headers)
        if re.fullmatch(r"topic|measurement|parameter|module type", header, re.I)
    ]
    specification_cols = [
        i for i, header in enumerate(headers)
        if re.search(r"description|customer\s+spec|requirement\s+text", header, re.I)
    ]
    result = []
    for row in rows[header_index + 1:] if len(rows) > 1 else rows:
        pairs = [
            f"{header}: {value}" if header and header != value else value
            for header, value in zip(headers, row) if value
        ]
        text = " | ".join(pairs)
        if not text or len(set(re.findall(r"\b(?:REQ|REF)-[\w.-]+", text))) > 1:
            continue
        measurement_text = "" if specification_cols else None
        if supplier_cols:
            measurement_text = "\n".join(
                row[i] or "" for i in context_cols + supplier_cols if i < len(row) and row[i]
            )
        result.append((text, measurement_text))
    return result


def _clean_evidence(text: str) -> str:
    lines = [
        line for line in text.splitlines()
        if not _DECORATION.search(line)
        and not re.fullmatch(r"\s*(?:page\s*)?\d{1,4}\s*", line, re.I)
    ]
    return _clean_block("\n".join(lines))


def _remove_repeated_edges(
    blocks: List[tuple[str, str, bool]],
) -> List[tuple[str, str]]:
    """Remove recurring page furniture only when it appears at the edge."""
    pages = {loc for loc, _, _ in blocks}
    counts = Counter(
        (loc, re.sub(r"\d+", "#", text.casefold().strip()))
        for loc, text, edge in blocks if edge and text.strip()
    )
    occurrences = Counter(key for _, key in counts)
    return [
        (loc, cleaned)
        for loc, text, edge in blocks
        if (cleaned := _clean_evidence(text))
        and not (edge and len(pages) >= 3 and
                 occurrences[re.sub(r"\d+", "#", text.casefold().strip())]
                 >= max(3, len(pages) // 2))
    ]


def _ocr_image(image: bytes, warnings: List[str]) -> str:
    if len(image) > _MAX_IMAGE_BYTES:
        warnings.append("An oversized evidence image was not OCR-processed")
        return ""
    try:
        from PIL import Image, UnidentifiedImageError
        import pytesseract
        from pytesseract import TesseractError, TesseractNotFoundError
        with Image.open(io.BytesIO(image)) as picture:
            if picture.width < 80 or picture.height < 40:
                return ""
            text = pytesseract.image_to_string(picture)
        return _clean_evidence(text)
    except ImportError:
        warnings.append("Image OCR unavailable: install Pillow and pytesseract")
    except (TesseractNotFoundError, TesseractError) as exc:
        warnings.append(f"Image OCR unavailable: {exc}")
    except (OSError, UnidentifiedImageError) as exc:
        warnings.append(f"Image OCR failed: {exc}")
    return ""


def _extract_pptx(content: bytes, file_name: str, warnings: List[str]) -> List[EvidenceChunk]:
    """
    Extract slide text from a PPTX (Office Open XML).

    Prefer python-pptx when installed (notes + tables). Fall back to a
    zip/XML reader so Azure Functions still work without the extra wheel.
    """
    try:
        from pptx import Presentation  # type: ignore
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = Presentation(io.BytesIO(content))
    except Exception as exc:
        warnings.append(f"PPTX structured parsing failed; using text-only XML: {exc}")
        return _extract_pptx_xml(content, file_name)

    blocks: List[tuple[str, str, bool, str]] = []
    supplier_texts: dict[tuple[str, str], str | None] = {}
    for i, slide in enumerate(prs.slides, 1):
        loc = f"Slide {i}"
        title = _clean_evidence(slide.shapes.title.text) if slide.shapes.title else ""
        image_count = 0

        def visit(shapes):
            nonlocal image_count
            for shape in shapes:
                edge = shape.top >= prs.slide_height * 0.88
                if shape.has_table:
                    rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                    for text, supplier_text in _table_row_blocks(rows):
                        blocks.append((loc, text, edge, "table_row"))
                        supplier_texts[(loc, _clean_evidence(text))] = supplier_text
                elif shape.has_text_frame:
                    blocks.append((loc, shape.text_frame.text, edge, "slide"))
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_count += 1
                    if image_count > _MAX_IMAGES_PER_PAGE:
                        warnings.append(f"{loc}: image OCR limit reached")
                        continue
                    ocr = _ocr_image(shape.image.blob, warnings)
                    if ocr:
                        blocks.append((loc, f"{title}\n{ocr}" if title else ocr, False, "image"))
                if shape.has_chart:
                    chart = shape.chart
                    chart_title = (
                        chart.chart_title.text_frame.text if chart.has_title else title
                    )
                    unit = re.search(
                        r"\((Hz|kHz|MHz|mA|A|V|ms|°C|%)\)",
                        chart_title, re.I,
                    )
                    for series in chart.series:
                        values = getattr(series, "values", ())
                        for index, value in enumerate(values):
                            if value is None:
                                continue
                            category = ""
                            try:
                                category = str(chart.plots[0].categories[index].label)
                            except (AttributeError, IndexError, TypeError):
                                pass
                            blocks.append((
                                loc,
                                f"{chart_title} {series.name} {category}: "
                                f"{value} {unit.group(1) if unit else ''}",
                                False,
                                "chart",
                            ))
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    visit(shape.shapes)

        visit(slide.shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            blocks.append((loc, slide.notes_slide.notes_text_frame.text, False, "notes"))

    kept = _remove_repeated_edges([(loc, text, edge) for loc, text, edge, _ in blocks])
    allowed = Counter(kept)
    chunks = []
    for loc, text, edge, kind in blocks:
        cleaned = _clean_evidence(text)
        if cleaned and allowed[(loc, cleaned)]:
            allowed[(loc, cleaned)] -= 1
            chunks.append(EvidenceChunk(
                file_name, loc, cleaned, len(chunks) + 1, kind,
                supplier_texts.get((loc, cleaned)),
            ))
    return chunks


def _extract_pptx_xml(content: bytes, file_name: str) -> List[EvidenceChunk]:
    """Minimal PPTX reader: concatenate every a:t run per slide XML part."""
    chunks: List[EvidenceChunk] = []
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        slide_names = sorted(
            n for n in zf.namelist()
            if re.match(r"ppt/slides/slide\d+\.xml$", n)
        )
        for name in slide_names:
            idx = int(re.search(r"slide(\d+)", name).group(1))
            xml = zf.read(name)
            root = ET.fromstring(xml)
            runs = [el.text for el in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t") if el.text]
            text = _clean_evidence("\n".join(runs))
            if text:
                chunks.append(EvidenceChunk(
                    file_name=file_name,
                    location=f"Slide {idx}",
                    text=text,
                    chunk_id=idx,
                    kind="slide",
                ))
    return chunks


def _extract_pdf(content: bytes, file_name: str, warnings: List[str]) -> List[EvidenceChunk]:
    try:
        import pdfplumber
        import pypdfium2
    except ImportError:
        warnings.append("PDF layout/OCR unavailable: install pdfplumber and pypdfium2")
        return _extract_pdf_text_only(content, file_name)

    blocks: List[tuple[str, str, bool, str]] = []
    supplier_texts: dict[tuple[str, str], str | None] = {}
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        renderer = pypdfium2.PdfDocument(content)
        try:
            for i, page in enumerate(pdf.pages, 1):
                loc = f"Page {i}"
                tables = page.find_tables()
                table_boxes = [
                    table.bbox for table in tables
                    if len(table.rows) > 1 and len(table.columns) > 1
                ]
                lines = page.extract_text_lines()
                for line in lines:
                    if any(
                        x0 <= line["x0"] and line["x1"] <= x1
                        and top <= line["top"] and line["bottom"] <= bottom
                        for x0, top, x1, bottom in table_boxes
                    ):
                        continue
                    blocks.append((
                        loc, line["text"],
                        line["top"] >= page.height * 0.88,
                        "page",
                    ))
                for table in tables:
                    for text, supplier_text in _table_row_blocks(table.extract()):
                        blocks.append((loc, text, False, "table_row"))
                        supplier_texts[(loc, _clean_evidence(text))] = supplier_text

                images = page.images
                if len(images) > _MAX_IMAGES_PER_PAGE:
                    warnings.append(f"{loc}: image OCR limit reached")
                needs_ocr = images[:_MAX_IMAGES_PER_PAGE] or not any(
                    l == loc and _clean_evidence(text) for l, text, _, _ in blocks
                )
                if needs_ocr:
                    rendered = renderer[i - 1].render(scale=2).to_pil()
                    for image in images[:_MAX_IMAGES_PER_PAGE]:
                        box = (
                            max(0, int(image["x0"] * 2)),
                            max(0, int(image["top"] * 2)),
                            min(rendered.width, int(image["x1"] * 2)),
                            min(rendered.height, int(image["bottom"] * 2)),
                        )
                        if box[2] <= box[0] or box[3] <= box[1]:
                            continue
                        data = io.BytesIO()
                        rendered.crop(box).save(data, format="PNG")
                        ocr = _ocr_image(data.getvalue(), warnings)
                        if ocr:
                            blocks.append((loc, ocr, False, "image"))
                    if not images:
                        data = io.BytesIO()
                        rendered.save(data, format="PNG")
                        ocr = _ocr_image(data.getvalue(), warnings)
                        if ocr:
                            blocks.append((loc, ocr, False, "image"))
        finally:
            renderer.close()

    kept = _remove_repeated_edges([(loc, text, edge) for loc, text, edge, _ in blocks])
    allowed = Counter(kept)
    chunks = []
    for loc, text, edge, kind in blocks:
        cleaned = _clean_evidence(text)
        if cleaned and allowed[(loc, cleaned)]:
            allowed[(loc, cleaned)] -= 1
            chunks.append(EvidenceChunk(
                file_name, loc, cleaned, len(chunks) + 1, kind,
                supplier_texts.get((loc, cleaned)),
            ))
    return chunks


def _extract_pdf_text_only(content: bytes, file_name: str) -> List[EvidenceChunk]:
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(content))
    except Exception:
        return []
    chunks: List[EvidenceChunk] = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = _clean_evidence(page.extract_text() or "")
        except Exception:
            text = ""
        if text:
            chunks.append(EvidenceChunk(
                file_name=file_name,
                location=f"Page {i}",
                text=text,
                chunk_id=i,
                kind="page",
            ))
    return chunks


def _extract_docx(content: bytes, file_name: str) -> List[EvidenceChunk]:
    return _extract_docx_xml(content, file_name)


def _extract_docx_xml(content: bytes, file_name: str) -> List[EvidenceChunk]:
    """Read paragraphs and supplier table cells without flattening the table."""
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    body = root.find(f"{ns}body")
    if body is None:
        raise ValueError("DOCX document has no body")

    def paragraph_text(para):
        return "".join(el.text or "" for el in para.iter(f"{ns}t"))

    chunks = []
    paragraph_count = table_count = 0
    for node in body:
        if node.tag == f"{ns}p":
            paragraph_count += 1
            text = _clean_evidence(paragraph_text(node))
            if text:
                chunks.append(EvidenceChunk(
                    file_name, f"Paragraph {paragraph_count}", text,
                    len(chunks) + 1, "paragraph",
                ))
        elif node.tag == f"{ns}tbl":
            table_count += 1
            rows = [
                [
                    "\n".join(paragraph_text(p) for p in cell.findall(f"{ns}p"))
                    for cell in row.findall(f"{ns}tc")
                ] for row in node.findall(f"{ns}tr")
            ]
            for row_index, (text, measured) in enumerate(_table_row_blocks(rows), 1):
                cleaned = _clean_evidence(text)
                if cleaned:
                    chunks.append(EvidenceChunk(
                        file_name, f"Table {table_count}, data row {row_index}", cleaned,
                        len(chunks) + 1, "table_row", measured,
                    ))
    return chunks


def _chunk_plain(text: str, file_name: str) -> List[EvidenceChunk]:
    """
    Split a TXT TDR on blank lines or 'Slide N' / 'Page N' markers so tests
    and exported-PPT text dumps keep location awareness.
    """
    text = text.replace("\r\n", "\n")
    marker = re.compile(
        r"^(?:---\s*)?(slide|page|section)\s+(\d+|[A-Za-z0-9._-]+)\s*:?\s*",
        re.I | re.M,
    )
    matches = list(marker.finditer(text))
    chunks: List[EvidenceChunk] = []
    if matches:
        for i, m in enumerate(matches):
            start = m.end()
            end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            body = _clean_evidence(text[start:end])
            if not body:
                continue
            kind = m.group(1).lower()
            loc = f"{kind.capitalize()} {m.group(2)}"
            chunks.append(EvidenceChunk(
                file_name=file_name, location=loc, text=body,
                chunk_id=i + 1, kind=kind if kind != "section" else "section",
            ))
        return chunks

    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    for i, block in enumerate(blocks, 1):
        cleaned = _clean_evidence(block)
        if not cleaned:
            continue
        chunks.append(EvidenceChunk(
            file_name=file_name,
            location=f"Block {i}",
            text=cleaned,
            chunk_id=i,
            kind="paragraph",
        ))
    return chunks


def _clean_block(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
